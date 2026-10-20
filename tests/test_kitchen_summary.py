"""廚房彙總: 取貨日 × 品項 box totals with a per-門市 breakdown, its .xlsx and the 總表's,
and the link to it at the end of every reply."""

from __future__ import annotations

import html
import io
import re
from datetime import date, datetime
from html.parser import HTMLParser

import pytest
from openpyxl import load_workbook

from bakery.pipeline import process_upload

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PENDING = "有待確認"


def _row(order_no, product, quantity, *, pickup_on, phone="0912-345-678", status=None) -> tuple:
    return (
        order_no, "王小姐", phone, date(2026, 10, 19), pickup_on, None,
        "陳志豪、林佳蓉", "美玲", product, quantity, None, status,
    )


TC01_ROWS = [
    _row("0301", "龍鳳喜餅", 100, pickup_on=date(2026, 10, 30)),
    _row("0301", "鳳梨酥禮盒", 50, pickup_on=date(2026, 10, 30), status="取消"),
    _row("0302", "龍鳳喜餅", 40, pickup_on=date(2026, 10, 30), phone="0922-111-333"),
    _row("0303", "龍鳳喜餅", 30, pickup_on=date(2026, 11, 10)),
]
CH01_ROWS = [
    _row("C-1", "龍鳳喜餅", 20, pickup_on=date(2026, 10, 30), phone="0933-444-555"),
    _row("C-1", "綠豆椪禮盒", 10, pickup_on=date(2026, 10, 30), phone="0933-444-555"),
    _row("C-2", "鳳梨酥禮盒", 60, pickup_on=date(2026, 11, 1), phone="0955-666-777"),
]
# 0302 again with another phone and more boxes: a 疑似單號重複使用, held for 總部.
TC01_REUSED = [_row("0302", "龍鳳喜餅", 90, pickup_on=date(2026, 10, 30), phone="0988-000-111")]

HEADER = ["取貨日", "星期", "門市", "龍鳳喜餅", "鳳梨酥禮盒", "綠豆椪禮盒", "合計"]


@pytest.fixture
def send(session, now, make_template_file):
    def send(store_code: str, rows) -> str:
        """`rows` for a 範本 of `store_code`, or the file's bytes as they are."""
        return process_upload(
            session,
            uploader_line_id=f"Udemo-{store_code}",
            filename=f"{store_code}.xlsx",
            file_bytes=rows if isinstance(rows, bytes) else make_template_file(store_code, rows),
            now=now,
        ).reply_text

    return send


@pytest.fixture
def orders(send):
    send("TC01", TC01_ROWS)
    send("CH01", CH01_ROWS)
    assert "待總部確認" in send("TC01", TC01_REUSED)


class _Tables(HTMLParser):
    """Each <table class="…">'s rows as cell texts (whitespace collapsed), by class."""

    def __init__(self) -> None:
        super().__init__()
        self.tables: dict[str, list[list[str]]] = {}
        self._table: list[list[str]] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._table = self.tables.setdefault(dict(attrs).get("class", ""), [])
        elif tag == "tr" and self._table is not None:
            self._table.append([])
        elif tag in ("td", "th") and self._table is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag == "table":
            self._table = None
        elif tag in ("td", "th") and self._cell is not None:
            self._table[-1].append(" ".join("".join(self._cell).split()))
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _tables(page: str) -> dict[str, list[list[str]]]:
    parser = _Tables()
    parser.feed(page)
    return parser.tables


def _sheets(content: bytes) -> dict[str, list[list[str]]]:
    """Every sheet's rows as the texts a reader sees: dates as yyyy/mm/dd, a 有待確認 comment
    after the number it is on."""

    def text(cell) -> str:
        v = cell.value
        if v is None:
            shown = ""
        elif isinstance(v, datetime):
            shown = f"{v:%Y/%m/%d}"
        else:
            shown = str(v)
        return f"{shown} {cell.comment.text}" if cell.comment else shown

    wb = load_workbook(io.BytesIO(content))
    return {ws.title: [[text(c) for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def _download(client, page: str, label: str):
    href = html.unescape(re.search(rf'<a href="([^"]+)">{label}</a>', page)[1])
    resp = client.get(href)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX
    return resp


def test_summary_adds_boxes_by_pickup_date_and_product_for_the_next_14_days(make_client, orders):
    page = make_client("kitchen").get("/kitchen").text

    assert "2026/10/21（三）～2026/11/03（二）" in page
    assert _tables(page)["summary"] == [
        HEADER,
        # 0301's 鳳梨酥禮盒 is cancelled; 0302 (held) still counts its 40 boxes.
        ["2026/10/30", "五", "全部門市", f"160 {PENDING}", "", "10", f"170 {PENDING}"],
        ["2026/10/30", "五", "彰化中華店", "20", "", "10", "30"],
        ["2026/10/30", "五", "台中中正店", f"140 {PENDING}", "", "", f"140 {PENDING}"],
        ["2026/11/01", "日", "全部門市", "", "60", "", "60"],
        ["2026/11/01", "日", "彰化中華店", "", "60", "", "60"],
    ]
    assert _tables(page)["pending"] == [["取貨日", "門市", "單號"], ["2026/10/30", "台中中正店", "0302"]]


def test_a_cancelled_line_back_in_production_counts_again(make_client, send, orders):
    send("TC01", [TC01_ROWS[0], _row("0301", "鳳梨酥禮盒", 50, pickup_on=date(2026, 10, 30))])

    rows = _tables(make_client("kitchen").get("/kitchen").text)["summary"]

    assert rows[1] == ["2026/10/30", "五", "全部門市", f"160 {PENDING}", "50", "10", f"220 {PENDING}"]


def test_an_applied_change_counts_with_its_new_quantity_unmarked(make_client, send, orders):
    send("CH01", [_row("C-1", "龍鳳喜餅", 25, pickup_on=date(2026, 10, 30), phone="0933-444-555"),
                  _row("C-1", "綠豆椪禮盒", 10, pickup_on=date(2026, 10, 30), phone="0933-444-555")])

    rows = _tables(make_client("kitchen").get("/kitchen").text)["summary"]

    assert rows[2] == ["2026/10/30", "五", "彰化中華店", "25", "", "10", "35"]


@pytest.mark.parametrize(
    ("params", "period", "dates"),
    [
        pytest.param({"start": "2026-11-01", "end": "2026-11-10"}, "2026/11/01（日）～2026/11/10（二）",
                     ["2026/11/01", "2026/11/10"], id="adjusted"),
        pytest.param({"start": "2026-11-10", "end": "2026-11-01"}, "2026/11/01（日）～2026/11/10（二）",
                     ["2026/11/01", "2026/11/10"], id="swapped"),
        pytest.param({"start": "2026-11-02", "end": "2026-11-02"}, "2026/11/02（一）～2026/11/02（一）",
                     [], id="one-empty-day"),
        pytest.param({"start": "昨天", "end": ""}, "2026/10/21（三）～2026/11/03（二）",
                     ["2026/10/30", "2026/11/01"], id="unreadable-is-default"),
    ],
)
def test_date_range_can_be_adjusted(make_client, orders, params, period, dates):
    page = make_client("kitchen").get("/kitchen", params=params).text

    assert period in page
    totals = [r[0] for r in _tables(page)["summary"][1:] if r[2] == "全部門市"]
    assert totals == dates
    if not dates:
        assert "這段期間沒有要做的喜餅" in page


def test_summary_download_matches_the_screen(make_client, orders):
    kitchen = make_client("kitchen")
    page = kitchen.get("/kitchen", params={"start": "2026-10-25", "end": "2026-11-15"}).text

    resp = _download(kitchen, page, "下載 .xlsx")

    sheets = _sheets(resp.content)
    tables = _tables(page)
    assert list(sheets) == ["廚房彙總", "有待確認的訂單"]
    assert sheets["廚房彙總"] == tables["summary"]
    assert sheets["有待確認的訂單"] == tables["pending"]
    assert "2026/11/10" in {r[0] for r in tables["summary"]}  # the range the page was showing
    assert "20261025-20261115" in resp.headers["content-disposition"]


def test_master_list_download_matches_the_screen(make_client, orders):
    hq = make_client("hq-lin")
    page = hq.get("/orders", params={"store": "TC01"}).text

    sheets = _sheets(_download(hq, page, "下載 .xlsx").content)

    assert list(sheets) == ["總表"]
    assert sheets["總表"] == _tables(page)["master-list"]
    assert len(sheets["總表"]) == 1 + 4  # header + TC01's lines, 取消 included
    assert ["2026/10/30", "台中中正店", "0301", "鳳梨酥禮盒", "50", "取消"] == sheets["總表"][2][:6]


def test_historical_orders_are_left_out(make_client, session, now, make_template_file, orders):
    process_upload(
        session,
        uploader_line_id=None,
        filename="CH02_歷史資料.xlsx",
        file_bytes=make_template_file("CH02", [_row("H-1", "龍鳳喜餅", 500, pickup_on=date(2026, 10, 30))]),
        now=now,
        historical=True,
    )

    rows = _tables(make_client("kitchen").get("/kitchen").text)["summary"]

    assert rows[1][3] == f"160 {PENDING}"
    assert "員林光明店" not in {r[2] for r in rows}


@pytest.mark.parametrize(
    ("store_code", "rows"),
    [
        pytest.param("TC01", TC01_ROWS, id="accepted"),
        pytest.param("TC01", b"not a workbook", id="refused"),
        pytest.param("TC01", [_row(f"9{i:03d}", f"不存在的品項{i}", 10, pickup_on=date(2026, 11, 28))
                              for i in range(400)], id="cut-to-the-line-limit"),
    ],
)
def test_every_reply_ends_with_the_kitchen_summary_link(send, monkeypatch, store_code, rows):
    monkeypatch.setenv("PUBLIC_BASE_URL", "https://bakery.example/")

    reply = send(store_code, rows)

    assert len(reply) <= 5000
    assert reply.endswith("\n\n廚房彙總：https://bakery.example/kitchen")
    assert reply.count("廚房彙總") == 1
