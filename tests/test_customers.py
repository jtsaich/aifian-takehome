"""客人資料: one 客人 per normalized phone, their order history, and the .xlsx export."""

from __future__ import annotations

import io
import json
import re
from datetime import date
from datetime import datetime as dt  # how openpyxl reads date cells back
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from bakery.models import Order
from bakery.pipeline import process_upload

MOCK = Path(__file__).resolve().parents[1] / "mock_data"
MANIFEST = json.loads((MOCK / "expectations.json").read_text())

# 單號, 訂購人姓名, 電話, 接單日, 取貨日, 宴客日, 新人姓名, 接單人, 品項, 盒數, 備註, 狀態
TC01_ROWS = (
    ("0201", "王小姐", "0912-345-678", date(2026, 10, 20), date(2026, 11, 28), None,
     None, None, "龍鳳喜餅", 120, None, None),
    ("0201", "王小姐", "0912-345-678", date(2026, 10, 20), date(2026, 11, 28), None,
     None, None, "鳳梨酥禮盒", 60, None, "取消"),
    ("0202", "李先生", "04-2345-6789", date(2026, 10, 21), date(2026, 12, 5), None,
     None, None, "綠豆椪禮盒", 40, None, None),
)
NT01_ROWS = (
    ("N001", "王美華", "+886 912 345 678", date(2026, 10, 10), date(2026, 12, 1), None,
     None, None, "綠豆椪禮盒", 40, None, None),
    ("N002", "林先生", "0912-3456", date(2026, 10, 12), date(2026, 12, 2), None,
     None, None, "龍鳳喜餅", 30, None, None),
    ("N003", "王小姐", "0912 345 678", date(2026, 10, 15), date(2026, 12, 3), None,
     None, None, "龍鳳喜餅", 50, None, "取消"),
)


@pytest.fixture
def two_stores(session, now, make_template_file) -> dict[str, int]:
    """Orders at TC01 and NT01; 單號 → Order.id."""
    for store, rows in (("TC01", TC01_ROWS), ("NT01", NT01_ROWS)):
        process_upload(
            session,
            uploader_line_id=f"Udemo-{store}",
            filename=f"{store}.xlsx",
            file_bytes=make_template_file(store, rows),
            now=now,
        )
    return {o.order_no: o.id for o in session.scalars(select(Order))}


def _customer_rows(html: str) -> dict[str, dict[str, str]]:
    """List page rows by phone: link, 訂單數, 總盒數 and the rest of the cells as text."""
    rows = {}
    for tr in re.findall(r'<tr class="customer">(.*?)</tr>', html, re.S):
        link, phone = re.search(r'<a href="([^"]+)">([^<]+)</a>', tr).groups()
        cells = [re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
        rows[phone] = {"link": link, "cells": cells}
    return rows


def _detail(client, phone: str) -> str:
    rows = _customer_rows(client.get("/customers").text)
    return client.get(rows[phone]["link"]).text


def test_every_way_of_writing_one_phone_is_one_customer(client, two_stores):
    rows = _customer_rows(client.get("/customers").text)

    assert set(rows) == {"0912345678", "0423456789"}
    assert rows["0912345678"]["cells"] == [
        "0912345678", "3", "160", "2026/10/10", "2026/10/20",
        "草屯和平店、台中中正店", "王美華、王小姐 姓名不一致",
    ]
    assert rows["0423456789"]["cells"] == [
        "0423456789", "1", "40", "2026/10/21", "2026/10/21", "台中中正店", "李先生",
    ]


def test_most_orders_come_first(client, two_stores):
    assert list(_customer_rows(client.get("/customers").text)) == ["0912345678", "0423456789"]


def test_customer_page_shows_the_history_and_links_every_order(client, two_stores):
    html = _detail(client, "0912345678")

    assert re.search(r'class="order-count">3<', html)
    assert re.search(r'class="total-boxes">160（不含取消）<', html)
    assert "2026/10/10" in html and "2026/10/20" in html
    assert "草屯和平店、台中中正店" in html
    assert "王美華、王小姐" in html
    assert "同一支電話留過不同的姓名" in html
    links = re.findall(r'<td class="order-no"><a href="(/orders/\d+)">([^<]+)</a>', html)
    assert links == [
        (f"/orders/{two_stores['N001']}", "N001"),
        (f"/orders/{two_stores['N003']}", "N003"),
        (f"/orders/{two_stores['0201']}", "0201"),
    ]
    assert "全部取消" in html  # N003


def test_one_name_is_not_flagged(client, two_stores):
    html = _detail(client, "0423456789")

    assert "李先生" in html
    assert "同一支電話留過不同的姓名" not in html


def test_any_order_of_a_customer_opens_the_same_page(client, two_stores):
    pages = {client.get(f"/customers/{two_stores[no]}").text for no in ("N001", "N003", "0201")}

    assert len(pages) == 1


def test_order_with_an_unusable_phone_belongs_to_no_customer_and_is_listed(client, two_stores):
    html = client.get("/customers").text

    unmatched = re.findall(r'<tr class="unmatched">(.*?)</tr>', html, re.S)
    assert len(unmatched) == 1
    assert f'href="/orders/{two_stores["N002"]}"' in unmatched[0] and "0912-3456" in unmatched[0]
    assert client.get(f"/customers/{two_stores['N002']}").status_code == 404


def test_phone_that_cannot_be_normalized_is_warned_about_on_upload(session, now, make_template_file):
    result = process_upload(
        session, uploader_line_id="Udemo-NT01", filename="NT01.xlsx",
        file_bytes=make_template_file("NT01", NT01_ROWS), now=now,
    )

    assert "電話「0912-3456」好像少了幾碼，已先收下，請確認號碼。" in result.reply_text
    assert "+886 912 345 678" not in result.reply_text


def test_unknown_order_is_not_found(client, two_stores):
    assert client.get("/customers/999999").status_code == 404


def test_phones_are_not_put_in_urls(client, two_stores):
    for page in ("/customers", f"/customers/{two_stores['0201']}"):
        hrefs = re.findall(r'href="([^"]*)"', client.get(page).text)
        assert hrefs
        assert not any(re.search(r"\d{7,}", h) for h in hrefs), hrefs


def test_customer_pages_offer_no_way_to_send_messages(client, two_stores):
    for page in ("/customers", f"/customers/{two_stores['0201']}"):
        html = client.get(page).text
        content = html.split("</nav>", 1)[1]
        assert "<form" not in content and "<button" not in content
        for scheme in ("tel:", "sms:", "mailto:", "line.me", "line://"):
            assert scheme not in html
        assert "不用來發訊息" in content


def test_export_has_the_same_customers_and_orders(client, two_stores):
    resp = client.get("/customers.xlsx")

    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    rows = lambda ws: [tuple(c.value for c in r) for r in ws.iter_rows()]  # noqa: E731
    assert wb.sheetnames == ["客人", "訂單"]
    assert rows(wb["客人"]) == [
        ("電話", "訂單數", "總盒數", "第一次接單日", "最近一次接單日", "下過單的門市", "訂購人姓名", "姓名不一致"),
        ("0912345678", 3, 160, dt(2026, 10, 10), dt(2026, 10, 20), "草屯和平店、台中中正店", "王美華、王小姐", "是"),
        ("0423456789", 1, 40, dt(2026, 10, 21), dt(2026, 10, 21), "台中中正店", "李先生", None),
    ]
    assert rows(wb["訂單"]) == [
        ("電話", "門市", "單號", "接單日", "取貨日", "訂購人姓名", "電話（原始寫法）", "盒數", "狀態"),
        ("0912345678", "草屯和平店", "N001", dt(2026, 10, 10), dt(2026, 12, 1), "王美華", "+886 912 345 678", 40, None),
        ("0912345678", "草屯和平店", "N003", dt(2026, 10, 15), dt(2026, 12, 3), "王小姐", "0912 345 678", 0, "全部取消"),
        ("0912345678", "台中中正店", "0201", dt(2026, 10, 20), dt(2026, 11, 28), "王小姐", "0912-345-678", 120, None),
        ("0423456789", "台中中正店", "0202", dt(2026, 10, 21), dt(2026, 12, 5), "李先生", "04-2345-6789", 40, None),
    ]


def test_export_without_orders_has_only_headers(client):
    wb = load_workbook(io.BytesIO(client.get("/customers.xlsx").content))

    assert wb["客人"].max_row == 1 and wb["訂單"].max_row == 1


DAY1_RETURNING = [
    rc for rc in MANIFEST["returning_customers"]
    if all(a["file"].startswith("day1/") for a in rc["appearances"])
]


@pytest.mark.parametrize("returning", DAY1_RETURNING, ids=[rc["phone_digits"] for rc in DAY1_RETURNING])
def test_day1_returning_customers_match_the_mock_data_table(client, session, now, returning):
    assert DAY1_RETURNING, "the 對照表 lists day-1 回頭客"
    for f in (f for f in MANIFEST["files"] if f["stage"] == "day1"):
        path = MOCK / f["file"]
        process_upload(
            session, uploader_line_id=f"Udemo-{f['store_code']}", filename=path.name,
            file_bytes=path.read_bytes(), now=now,
        )

    html = _detail(client, returning["phone_digits"])

    order_nos = re.findall(r'<td class="order-no"><a href="/orders/\d+">([^<]+)</a>', html)
    assert sorted(order_nos) == sorted(a["ref"].removeprefix("單號 ") for a in returning["appearances"])
    names = {a["orderer_name"] for a in returning["appearances"]}
    for name in names:
        assert name in html
    assert ("同一支電話留過不同的姓名" in html) is (len(names) > 1)
