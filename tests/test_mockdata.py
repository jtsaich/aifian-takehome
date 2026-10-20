"""Mock data: determinism, seeded errors present in the actual files, 對照表 consistent with them.

Checks read the .xlsx files directly instead of trusting the generator's own
bookkeeping, so a wrong row number or a missing seeded error fails here.
"""

from __future__ import annotations

import json
import re
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bakery.catalog import PRODUCTS, STORES
from bakery.mockdata import DEFAULT_SEED, generate_mock_data
from bakery.mockdata.__main__ import main
from bakery.mockdata.expectations import (
    REQUIRED_SEEDED,
    ChangeKind,
    FindingKind,
    Outcome,
    Seeded,
)
from bakery.template import COLUMNS, ORDER_COLUMNS

TODAY = date(2026, 10, 21)
REPO_ROOT = Path(__file__).resolve().parents[1]
ROC = re.compile(r"^(1\d{2})/(\d{2})/(\d{2})$")
HEADER_ROW = 3
ALIASES = {a for p in PRODUCTS for a in p.aliases}
CANONICAL = {p.name for p in PRODUCTS}
# Header synonyms are tolerated, not "extra" columns.
SYNONYMS = {"手機": "電話"}


def _files(root: Path) -> dict[str, bytes]:
    return {str(p.relative_to(root)): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()}


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and (m := ROC.match(value)):
        return date(int(m[1]) + 1911, int(m[2]), int(m[3]))
    return None


def _rows(path: Path, sheet: str | None = None) -> tuple[list[str], dict[int, dict[str, object]]]:
    """(headers, {excel_row: {field: value}}) for a template-format sheet."""
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    headers = [c.value for c in ws[HEADER_ROW] if c.value is not None]
    fields = [SYNONYMS.get(h, h) for h in headers]
    rows = {}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        values = {f: c.value for f, c in zip(fields, row)}
        if any(v is not None for v in values.values()):
            rows[row[0].row] = values
    return headers, rows


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    out = tmp_path_factory.mktemp("mock")
    return out, generate_mock_data(out, seed=DEFAULT_SEED, today=TODAY)


def _day1_paths(out: Path, manifest) -> list[Path]:
    return [out / f.file for f in manifest.files if f.stage == "day1"]


# ---------------------------------------------------------------- determinism


def test_same_seed_and_today_produce_byte_identical_files(tmp_path):
    generate_mock_data(tmp_path / "a", seed=7, today=TODAY)
    generate_mock_data(tmp_path / "b", seed=7, today=TODAY)
    a, b = _files(tmp_path / "a"), _files(tmp_path / "b")
    assert a.keys() == b.keys()
    assert [k for k in a if a[k] != b[k]] == []


def test_different_seed_changes_the_fake_people(tmp_path):
    generate_mock_data(tmp_path / "a", seed=7, today=TODAY)
    generate_mock_data(tmp_path / "b", seed=8, today=TODAY)
    a, b = _files(tmp_path / "a"), _files(tmp_path / "b")
    assert a["day1/TC01_台中中正店.xlsx"] != b["day1/TC01_台中中正店.xlsx"]


@pytest.mark.parametrize("today", [date(2026, 10, 21), date(2027, 3, 5)])
def test_dates_are_relative_to_today(tmp_path, today):
    manifest = generate_mock_data(tmp_path, seed=DEFAULT_SEED, today=today)
    past_rows = {
        (f.file, r) for f in manifest.files if f.stage == "day1" for fd in f.findings
        if fd.kind == FindingKind.PICKUP_IN_PAST for r in fd.rows
    }
    for path in _day1_paths(tmp_path, manifest):
        _, rows = _rows(path)
        for r, v in rows.items():
            pickup = _as_date(v["取貨日"])
            if pickup is None:
                continue  # the seeded unparseable date
            rel = str(path.relative_to(tmp_path))
            if (rel, r) in past_rows:
                assert today - timedelta(days=30) <= pickup < today
            else:
                assert today <= pickup <= today + timedelta(days=90), (rel, r, pickup)
    for lg in manifest.legacy:
        for o in lg.orders:
            assert date.fromisoformat(o.pickup_on) < today


# ---------------------------------------------------------------- seeded errors


def _present_issues(path: Path, today: date) -> set[Seeded]:
    """Which required issues are physically present in a store file."""
    found: set[Seeded] = set()
    wb = load_workbook(path)
    if "備忘" in wb.sheetnames:
        found.add(Seeded.MEMO_SHEET)
    headers, rows = _rows(path)
    if any(h not in COLUMNS and h not in SYNONYMS for h in headers):
        found.add(Seeded.EXTRA_COLUMN)
    by_order: dict[object, list[dict]] = {}
    for v in rows.values():
        if v.get("單號") in (None, ""):
            found.add(Seeded.MISSING_ORDER_NO)
        else:
            by_order.setdefault(v["單號"], []).append(v)
        if any(isinstance(v.get(c), str) and ROC.match(v[c]) for c in ("接單日", "取貨日", "宴客日")):
            found.add(Seeded.ROC_DATE)
        if v.get("品項") in ALIASES:
            found.add(Seeded.PRODUCT_ALIAS)
        pickup = _as_date(v.get("取貨日"))
        if pickup is not None and pickup < today:
            found.add(Seeded.PICKUP_IN_PAST)
    for lines in by_order.values():
        if any(len({repr(line.get(c)) for line in lines}) > 1 for c in ORDER_COLUMNS):
            found.add(Seeded.INCONSISTENT_ORDER_FIELDS)
    return found


@pytest.mark.parametrize("seed", [1, DEFAULT_SEED, 9999])
def test_every_required_error_type_is_in_at_least_one_store_file(tmp_path, seed):
    manifest = generate_mock_data(tmp_path, seed=seed, today=TODAY)
    day1 = [f for f in manifest.files if f.stage == "day1"]
    assert sorted(f.store_code for f in day1) == sorted(s.code for s in STORES)
    present = {f.file: _present_issues(tmp_path / f.file, TODAY) for f in day1}
    for f in day1:
        assert present[f.file], f"{f.file} has no seeded error"
        # The 對照表 claims exactly what the file contains (for the required kinds).
        assert present[f.file] == {s for s in f.seeded if s in REQUIRED_SEEDED}, f.file
    assert set().union(*present.values()) == set(REQUIRED_SEEDED)


def test_findings_point_at_the_right_rows(generated):
    out, manifest = generated
    for f in (f for f in manifest.files if f.stage == "day1"):
        _, rows = _rows(out / f.file)
        for fd in f.findings:
            if fd.kind == FindingKind.SKIPPED_SHEET:
                assert fd.sheet in load_workbook(out / f.file).sheetnames
                continue
            cells = [rows[r] for r in fd.rows]
            for v in cells:
                if fd.order_nos:
                    assert v["單號"] in fd.order_nos
            match fd.kind:
                case FindingKind.MISSING_ORDER_NO:
                    assert all(v["單號"] is None for v in cells)
                case FindingKind.INVALID_QUANTITY:
                    assert all(not (isinstance(v["盒數"], int) and v["盒數"] > 0) for v in cells)
                case FindingKind.UNKNOWN_PRODUCT:
                    assert all(v["品項"] not in CANONICAL | ALIASES for v in cells)
                case FindingKind.UNPARSEABLE_DATE:
                    assert all(_as_date(v["取貨日"]) is None for v in cells)
                case FindingKind.INCONSISTENT_ORDER_FIELDS:
                    assert len({repr(v["取貨日"]) for v in cells}) > 1
                case FindingKind.PICKUP_IN_PAST:
                    assert all(_as_date(v["取貨日"]) < TODAY for v in cells)
                case FindingKind.PICKUP_TOO_SOON:
                    assert all(0 <= (_as_date(v["取貨日"]) - TODAY).days <= 3 for v in cells)
                case FindingKind.QUANTITY_OVER_THRESHOLD:
                    assert all(v["盒數"] > 300 for v in cells)
                case FindingKind.PHONE_FORMAT:
                    assert all(len(re.sub(r"\D", "", v["電話"])) < 10 for v in cells)
                case FindingKind.DUPLICATE_PHONE_PICKUP:
                    assert len({(v["電話"], repr(v["取貨日"])) for v in cells}) == 1
                    assert len({v["單號"] for v in cells}) > 1
        accepted_rows = [r for r in rows if r not in f.rejected_rows]
        assert len(accepted_rows) == f.accepted_lines
        assert {rows[r]["單號"] for r in accepted_rows} == set(f.accepted_orders)


def test_alias_rows_in_the_table_use_aliases_in_the_file(generated):
    out, manifest = generated
    f = next(f for f in manifest.files if Seeded.PRODUCT_ALIAS in f.seeded)
    _, rows = _rows(out / f.file)
    alias_tolerances = [t for t in f.tolerances if t.kind == "product_alias"]
    assert alias_tolerances
    for tol in alias_tolerances:
        (r,) = tol.rows
        assert rows[r]["品項"] == tol.detail["寫法"]
        product = next(p for p in PRODUCTS if p.name == tol.detail["應解析為"])
        assert rows[r]["品項"] in product.aliases


# ---------------------------------------------------------------- day 2


def _orders(path: Path) -> dict[str, list[dict]]:
    _, rows = _rows(path)
    out: dict[str, list[dict]] = {}
    for v in rows.values():
        out.setdefault(v["單號"], []).append(v)
    return out


def test_day2_covers_every_change_kind_and_cutoff(generated):
    out, manifest = generated
    changed = next(f for f in manifest.files if f.stage == "day2" and f.changes)
    before = _orders(out / changed.process_after[0])
    after = _orders(out / changed.file)
    day2 = date.fromisoformat(manifest.day2)
    assert day2 == TODAY + timedelta(days=1)

    kinds = {k for c in changed.changes for k in c.kinds}
    assert kinds == set(ChangeKind)
    for c in changed.changes:
        assert before[c.order_no] != after[c.order_no], c.order_no
        old_pickup = _as_date(before[c.order_no][0]["取貨日"])
        assert c.past_cutoff == ((old_pickup - day2).days < manifest.assumptions["cutoff_days"])
        production = set(c.kinds) - {ChangeKind.INFO_FIX}
        if ChangeKind.SUSPECTED_REUSE in c.kinds or (c.past_cutoff and production):
            assert c.outcome == Outcome.PENDING
        else:
            assert c.outcome == Outcome.APPLIED
        if ChangeKind.CANCEL in c.kinds:
            assert all(line["狀態"] == "取消" for line in after[c.order_no])
        if ChangeKind.SUSPECTED_REUSE in c.kinds:
            assert before[c.order_no][0]["電話"] != after[c.order_no][0]["電話"]
    # A past-cutoff change of each class: one held for review, one 資料更正 applied.
    past = [c for c in changed.changes if c.past_cutoff]
    assert {c.outcome for c in past} == {Outcome.PENDING, Outcome.APPLIED}

    for no in changed.unchanged_orders:
        assert before[no] == after[no]
    for no in changed.absent_orders:
        assert no in before and no not in after
    for no in changed.new_orders:
        assert no not in before and no in after


def test_day2_has_an_identical_full_resend(generated):
    out, manifest = generated
    resend = next(f for f in manifest.files if f.stage == "day2" and not f.changes and not f.new_orders)
    assert (out / resend.file).read_bytes() == (out / resend.process_after[0]).read_bytes()
    assert set(resend.unchanged_orders) == set(_orders(out / resend.file))


# ---------------------------------------------------------------- legacy & templates


def test_two_legacy_files_have_a_different_layout_and_no_order_no(generated):
    out, manifest = generated
    assert len(manifest.legacy) == 2
    assert len({lg.store_code for lg in manifest.legacy}) == 2
    for lg in manifest.legacy:
        wb = load_workbook(out / lg.file)
        for ws in wb.worksheets:
            headers = [c.value for c in ws[1]]
            assert "單號" not in headers
            assert headers != list(COLUMNS)
        for o in lg.orders:
            ws = wb[o.source_sheet]
            for r in o.source_rows:
                assert o.orderer_name in [c.value for c in ws[r]]


def test_mock_files_keep_the_catalog_dropdown_even_with_reordered_columns(generated):
    out, manifest = generated
    f = next(f for f in manifest.files if Seeded.REORDERED_COLUMNS in f.seeded)
    ws = load_workbook(out / f.file).worksheets[0]
    col = next(c for c in ws[HEADER_ROW] if c.value == "品項").column_letter
    dv = next(dv for dv in ws.data_validations.dataValidation if f"{col}4" in dv.sqref)
    assert dv.formula1.strip('"').split(",") == [p.name for p in PRODUCTS]


# ---------------------------------------------------------------- one command


def test_one_command_regenerates_everything(tmp_path, capsys):
    assert main(["--out", str(tmp_path), "--today", "2026-10-21"]) == 0
    first = _files(tmp_path)
    assert "對照表.md" in first and "expectations.json" in first
    data = json.loads(first["expectations.json"])
    listed = data["templates"] + [f["file"] for f in data["files"]] + [lg["file"] for lg in data["legacy"]]
    assert sorted(listed) == sorted(k for k in first if k.endswith(".xlsx"))
    assert main(["--out", str(tmp_path), "--today", "2026-10-21"]) == 0
    assert _files(tmp_path) == first


def _members(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_committed_mock_data_matches_the_generator(tmp_path):
    """mock_data/ in the repo is exactly what the generator produces for its recorded seed and day 1."""
    committed = REPO_ROOT / "mock_data"
    recorded = json.loads((committed / "expectations.json").read_text(encoding="utf-8"))
    generate_mock_data(tmp_path, seed=recorded["seed"], today=date.fromisoformat(recorded["today"]))
    ours, theirs = _files(tmp_path), _files(committed)
    assert ours.keys() == theirs.keys()
    for name in ours:
        if name.endswith(".xlsx"):
            # Compare contents, not deflate output, which can differ between zlib builds.
            assert _members(tmp_path / name) == _members(committed / name), name
        else:
            assert ours[name] == theirs[name], name
