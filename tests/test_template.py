import io
from pathlib import Path

from openpyxl import load_workbook

from bakery.catalog import PRODUCTS, STORES, product_names
from bakery.parsing import is_template_file
from bakery.template import write_template


def _write(tmp_path: Path, store) -> Path:
    return write_template(tmp_path / "t.xlsx", store, product_names())


EXPECTED_HEADERS = [
    "單號",
    "訂購人姓名",
    "電話",
    "接單日",
    "取貨日",
    "宴客日",
    "新人姓名",
    "接單人",
    "品項",
    "盒數",
    "備註",
    "狀態",
]


def _list_options(ws, header: str) -> list[str]:
    """Options of the list dropdown covering the first data cell under `header`."""
    header_cells = {c.value: c for c in ws[3]}
    target = ws.cell(row=4, column=header_cells[header].column).coordinate
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and target in dv.sqref:
            return dv.formula1.strip('"').split(",")
    raise AssertionError(f"no list dropdown on {header} ({target})")


def test_template_has_one_header_row_with_order_and_line_columns(tmp_path: Path):
    path = _write(tmp_path, STORES[0])
    ws = load_workbook(path).active
    assert [c.value for c in ws[3]] == EXPECTED_HEADERS


def test_template_puts_store_code_and_version_in_fixed_cells(tmp_path: Path):
    store = STORES[2]
    ws = load_workbook(_write(tmp_path, store)).active
    assert ws["A1"].value == "門市代碼"
    assert ws["B1"].value == store.code
    assert ws["E1"].value == "範本版本"
    assert ws["F1"].value == "v1"


def test_product_dropdown_lists_exactly_the_catalog_names(tmp_path: Path):
    ws = load_workbook(_write(tmp_path, STORES[0])).active
    assert _list_options(ws, "品項") == [
        "龍鳳喜餅",
        "鳳梨酥禮盒",
        "綠豆椪禮盒",
        "太陽餅禮盒",
        "咖哩大餅",
        "綜合西式喜餅",
        "蛋黃酥禮盒",
    ]
    assert _list_options(ws, "品項") == [p.name for p in PRODUCTS]


def test_status_dropdown_only_offers_cancel(tmp_path: Path):
    ws = load_workbook(_write(tmp_path, STORES[0])).active
    assert _list_options(ws, "狀態") == ["取消"]


def test_hq_downloads_a_stores_template_built_from_the_live_product_list(client, session):
    client.post("/products", data={"kind": "product", "name": "芋頭酥禮盒"})

    resp = client.get("/templates/TC02.xlsx")

    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert (ws["B1"].value, ws["D1"].value) == ("TC02", "台中北屯店")
    assert _list_options(ws, "品項") == [*(p.name for p in PRODUCTS), "芋頭酥禮盒"]
    assert is_template_file(resp.content)
    assert '<a href="/templates/TC02.xlsx">' in client.get("/products").text
    assert client.get("/templates/XX99.xlsx").status_code == 404
