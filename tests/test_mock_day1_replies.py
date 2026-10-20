"""Every day-1 mock file through process_upload: what is kept, what is refused, and
what the reply says must match the 對照表 (mock_data/expectations.json)."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from bakery.models import Order
from bakery.pipeline import process_upload

MOCK = Path(__file__).resolve().parents[1] / "mock_data"
MANIFEST = json.loads((MOCK / "expectations.json").read_text())
DAY1 = [f for f in MANIFEST["files"] if f["stage"] == "day1"]


def _order_no_at(path: Path, sheet: str, row: int) -> str:
    ws = load_workbook(path, data_only=True)[sheet]
    headers = [c.value for c in ws[3]]
    return str(ws.cell(row=row, column=headers.index("單號") + 1).value)


@pytest.fixture(params=DAY1, ids=[f["file"] for f in DAY1])
def processed(request, session, now):
    expected = request.param
    assert expected["upload_on"] == now.date().isoformat()
    path = MOCK / expected["file"]
    result = process_upload(
        session,
        uploader_line_id=f"Udemo-{expected['store_code']}",
        filename=path.name,
        file_bytes=path.read_bytes(),
        now=now,
    )
    orders = {o.order_no: o for o in session.scalars(select(Order))}
    return expected, path, result, orders


def test_accepted_orders_and_lines_match(processed):
    expected, _, result, orders = processed

    assert sorted(orders) == sorted(expected["accepted_orders"])
    assert sum(len(o.lines) for o in orders.values()) == expected["accepted_lines"]
    assert result.outcome.order_count == len(expected["accepted_orders"])


def test_rejected_rows_match(processed):
    expected, _, result, _ = processed

    rejected = [(loc.sheet, row) for loc in result.outcome.rejected_rows for row in loc.rows]
    assert rejected == [("喜餅訂單", row) for row in expected["rejected_rows"]]


def test_findings_match_kind_place_and_order(processed):
    expected, _, result, _ = processed

    actual = [
        (f.kind.value, f.severity.value, [(loc.sheet, list(loc.rows)) for loc in f.locations], list(f.order_nos))
        for f in result.outcome.findings
    ]
    assert actual == [
        (f["kind"], f["severity"], [(f["sheet"], f["rows"])], f["order_nos"]) for f in expected["findings"]
    ]


def test_reply_has_the_summary_and_every_problem_in_store_wording(processed):
    expected, _, result, _ = processed

    assert result.reply_text.startswith(expected["reply_summary"])
    for finding in expected["findings"]:
        assert finding["reply"] in result.reply_text


def test_tolerated_deviations_are_read_correctly_and_not_reported(processed):
    expected, path, _, orders = processed

    for tol in expected["tolerances"]:
        detail = tol["detail"]
        if tol["kind"] == "roc_date":
            for row in tol["rows"]:
                o = orders[_order_no_at(path, tol["sheet"], row)]
                assert o.received_on == date.fromisoformat(detail["接單日應解析為"])
                assert o.pickup_on == date.fromisoformat(detail["取貨日應解析為"])
        elif tol["kind"] == "product_alias":
            o = orders[_order_no_at(path, tol["sheet"], tol["rows"][0])]
            assert detail["應解析為"] in {line.product.name for line in o.lines}
        elif tol["kind"] == "header_synonym":
            assert all(o.phone for o in orders.values())
        else:
            assert tol["kind"] in {"reordered_columns", "extra_column"}  # covered by the counts above
