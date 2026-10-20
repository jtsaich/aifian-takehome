"""Read a 範本 .xlsx into 喜餅訂單 / 訂單明細, grouped by 單號.

Pure: bytes in, dataclasses out; no database. The 品項 list comes in as a
lookup so it can follow what HQ maintains.

Current scope is a clean template file: the 喜餅訂單 sheet, exact template
headers, Excel dates, known 品項 names or 品項別名.
"""

from __future__ import annotations

import io
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from bakery import template as t

# 品項 name or 品項別名 -> product code.
ProductLookup = Mapping[str, str]


@dataclass(frozen=True)
class ParsedLine:
    """訂單明細: one template row."""

    sheet: str
    row: int  # Excel row number, for pointing the 門市 at it
    product_code: str
    quantity: int
    note: str | None
    cancelled: bool


@dataclass(frozen=True)
class ParsedOrder:
    """喜餅訂單: the order-level columns (taken from its first row) and its rows."""

    order_no: str
    orderer_name: str | None
    phone: str | None
    received_on: date
    pickup_on: date
    banquet_on: date | None
    couple_names: str | None
    taken_by: str | None
    lines: tuple[ParsedLine, ...]


@dataclass(frozen=True)
class ParsedFile:
    store_code: str
    template_version: str | None
    orders: tuple[ParsedOrder, ...]

    @property
    def line_count(self) -> int:
        return sum(len(o.lines) for o in self.orders)


def parse_upload(file_bytes: bytes, products: ProductLookup) -> ParsedFile:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb[t.SHEET_NAME] if t.SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    store_code = _text(ws[t.STORE_CODE_CELL].value)
    if store_code is None:
        raise ValueError(f"no 門市代碼 in {t.STORE_CODE_CELL}")
    return ParsedFile(
        store_code=store_code,
        template_version=_text(ws[t.VERSION_CELL].value),
        orders=_orders(ws, products),
    )


def _orders(ws: Worksheet, products: ProductLookup) -> tuple[ParsedOrder, ...]:
    columns = {
        name: idx for idx, cell in enumerate(ws[t.HEADER_ROW]) if (name := _text(cell.value)) is not None
    }
    rows_by_order: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    for row_no, values in enumerate(
        ws.iter_rows(min_row=t.FIRST_DATA_ROW, values_only=True), start=t.FIRST_DATA_ROW
    ):
        if all(_text(v) is None for v in values):
            continue
        cells = {name: values[idx] if idx < len(values) else None for name, idx in columns.items()}
        order_no = _text(cells[t.COL_ORDER_NO])
        if order_no is None:
            raise ValueError(f"row {row_no}: no 單號")
        rows_by_order.setdefault(order_no, []).append((row_no, cells))

    return tuple(
        _order(ws.title, order_no, rows, products) for order_no, rows in rows_by_order.items()
    )


def _order(
    sheet: str, order_no: str, rows: list[tuple[int, dict[str, Any]]], products: ProductLookup
) -> ParsedOrder:
    first_row, first = rows[0]
    return ParsedOrder(
        order_no=order_no,
        orderer_name=_text(first[t.COL_ORDERER_NAME]),
        phone=_text(first[t.COL_PHONE]),
        received_on=_required_date(first[t.COL_RECEIVED_ON], first_row),
        pickup_on=_required_date(first[t.COL_PICKUP_ON], first_row),
        banquet_on=_date(first[t.COL_BANQUET_ON], first_row),
        couple_names=_text(first[t.COL_COUPLE_NAMES]),
        taken_by=_text(first[t.COL_TAKEN_BY]),
        lines=tuple(
            ParsedLine(
                sheet=sheet,
                row=row_no,
                product_code=products[_text(cells[t.COL_PRODUCT]) or ""],
                quantity=int(cells[t.COL_QUANTITY]),
                note=_text(cells[t.COL_NOTE]),
                cancelled=_text(cells[t.COL_STATUS]) == t.STATUS_CANCELLED,
            )
            for row_no, cells in rows
        ),
    )


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _date(value: Any, row: int) -> date | None:
    if value is None or _text(value) is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"row {row}: not an Excel date")


def _required_date(value: Any, row: int) -> date:
    parsed = _date(value, row)
    if parsed is None:
        raise ValueError(f"row {row}: missing date")
    return parsed
