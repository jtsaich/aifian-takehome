"""`bakery-demo`: replay the nine-step demo against a running site, over HTTP.

    uv run bakery-demo --reset                       # local `docker compose up`
    uv run bakery-demo --base-url https://<service>  # a deployed site with an empty database

It regenerates the mock files with today's date (so the cutoff case is past the cutoff
on any run date), logs in as a demo 總部人員, sends each file through 模擬群組上傳 (the
same `process_upload` as the LINE group), and prints every reply and the page to open.

`--reset` empties the database first. Locally it runs `bakery-init --wipe-all-data` in
the compose `init` service; for any other site it needs `--database-url` and the
explicit `--wipe-remote`, so a deployed database is never wiped by accident.
"""

from __future__ import annotations

import argparse
import html
import io
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from enum import Enum
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from openpyxl import Workbook, load_workbook
from sqlalchemy.engine import make_url

from bakery import legacy, template as t
from bakery.catalog import store_by_code
from bakery.clock import TAIPEI, today_taipei
from bakery.db import init_database
from bakery.mockdata import DEFAULT_SEED, generate_mock_data
from bakery.seed import DEMO_PASSWORD
from bakery.web.deps import XLSX_MEDIA_TYPE
from bakery.xlsx import save_workbook

DEFAULT_BASE_URL = "http://localhost:8000"
DEFAULT_OUT_DIR = Path("demo_data")
DEMO_USERNAME = "hq-lin"
LOCAL_HOSTS = frozenset({"localhost", "127.0.0.1", "::1"})
# The kitchen totals the demo prints cover this many days from today.
KITCHEN_DAYS = 90
COMPOSE_RESET = ("docker", "compose", "run", "--rm", "init", "bakery-init", "--wipe-all-data")

# The TC01 orders the story edits (see mock_data/對照表.md).
FIXED_ORDER_NO = "0113"  # the row day 1 sent without a 單號
ZERO_ORDER_NO = "0112"  # 盒數 0 on day 1
FIXED_ZERO_QUANTITY = 60
ADD_ORDER_NO, ADD_PRODUCT, ADD_FROM, ADD_TO = "0101", "龍鳳喜餅", 80, 120
# 0106 is picked up 6 days after the generator's day 1 (today), inside the 7-day cutoff.
LATE_ORDER_NO, LATE_PRODUCT, LATE_FROM, LATE_TO = "0106", "綠豆椪禮盒", 50, 70


class DemoError(Exception):
    """The demo cannot go on; the message says why, in 繁中."""


class Reset(Enum):
    COMPOSE = "compose"  # `bakery-init --wipe-all-data` in the local compose `init` service
    DATABASE_URL = "database_url"  # wipe and seed in-process through --database-url


def is_local(base_url: str) -> bool:
    return (urlsplit(base_url).hostname or "") in LOCAL_HOSTS


def reset_plan(base_url: str, database_url: str | None, *, wipe_remote: bool) -> Reset:
    """How `--reset` empties the database; refuses to touch a non-local one without
    `--wipe-remote`."""
    if database_url is None:
        if not is_local(base_url):
            raise DemoError(
                f"{base_url} 不是本機。要清空部署環境的資料庫，請同時給 --database-url 和 --wipe-remote。"
            )
        return Reset.COMPOSE
    local = is_local(base_url) and (make_url(database_url).host or "") in LOCAL_HOSTS
    if not local and not wipe_remote:
        raise DemoError(
            "--database-url 指向的不是本機資料庫：清空會刪掉上面所有的訂單和上傳紀錄。"
            "確定要清空，請再加 --wipe-remote。"
        )
    return Reset.DATABASE_URL


def reset_database(plan: Reset, database_url: str | None) -> None:
    if plan is Reset.DATABASE_URL:
        assert database_url is not None
        url = make_url(database_url).set(drivername="postgresql+psycopg")
        init_database(url.render_as_string(hide_password=False), wipe=True)
        return
    done = subprocess.run(COMPOSE_RESET, capture_output=True, text=True)
    if done.returncode != 0:
        raise DemoError(
            f"清空本機資料庫失敗（{' '.join(COMPOSE_RESET)}），請在 repo 根目錄執行，"
            f"並先 docker compose up --build：\n{done.stderr.strip()}"
        )


# ---------------------------------------------------------------- files


@dataclass(frozen=True)
class DemoFiles:
    with_errors: Path  # step 1
    fixed: Path  # step 2
    added: Path  # step 3
    late_change: Path  # steps 4 and 6
    reverted: Path  # step 7
    other_stores: tuple[Path, ...]  # step 8: the other 門市's day-1 files
    historical: tuple[Path, ...]  # step 9: converted legacy files


def prepare_files(out_dir: Path, today: date, now: datetime) -> DemoFiles:
    """Mock data for `today`, plus the TC01 files the story sends, under `out_dir`."""
    shutil.rmtree(out_dir, ignore_errors=True)
    manifest = generate_mock_data(out_dir, seed=DEFAULT_SEED, today=today)
    day1 = {f.store_code: out_dir / f.file for f in manifest.files if f.stage == "day1"}
    steps = out_dir / "demo"

    wb = load_workbook(day1["TC01"])
    with_errors = save_workbook(wb, steps / "1_TC01_台中中正店_有錯.xlsx", now)
    sheet = _Sheet(wb)
    sheet.set(None, None, t.COL_ORDER_NO, FIXED_ORDER_NO)
    sheet.set(ZERO_ORDER_NO, None, t.COL_QUANTITY, FIXED_ZERO_QUANTITY)
    fixed = save_workbook(wb, steps / "2_TC01_台中中正店_修正.xlsx", now)
    sheet.set(ADD_ORDER_NO, ADD_PRODUCT, t.COL_QUANTITY, ADD_TO)
    added = save_workbook(wb, steps / "3_TC01_台中中正店_追加.xlsx", now)
    sheet.set(LATE_ORDER_NO, LATE_PRODUCT, t.COL_QUANTITY, LATE_TO)
    late_change = save_workbook(wb, steps / "4_TC01_台中中正店_過截止點追加.xlsx", now)
    sheet.set(LATE_ORDER_NO, LATE_PRODUCT, t.COL_QUANTITY, LATE_FROM)
    reverted = save_workbook(wb, steps / "7_TC01_台中中正店_改回原本.xlsx", now)

    historical = []
    for item in manifest.legacy:
        store = store_by_code(item.store_code)
        converted = legacy.convert(store.code, (out_dir / item.file).read_bytes())
        path = steps / f"9_{store.code}_{store.name}_歷史資料.xlsx"
        path.write_bytes(converted.content)
        historical.append(path)

    return DemoFiles(
        with_errors=with_errors,
        fixed=fixed,
        added=added,
        late_change=late_change,
        reverted=reverted,
        other_stores=tuple(path for code, path in day1.items() if code != "TC01"),
        historical=tuple(historical),
    )


class _Sheet:
    """Cell edits on a 範本 by 單號 (None: the row without one) and 品項."""

    def __init__(self, wb: Workbook) -> None:
        self.ws = wb[t.SHEET_NAME]
        header = next(self.ws.iter_rows(min_row=t.HEADER_ROW, max_row=t.HEADER_ROW, values_only=True))
        self.col = {name: i for i, name in enumerate(header, start=1) if name}

    def set(self, order_no: str | None, product: str | None, column: str, value) -> None:
        for row in range(t.FIRST_DATA_ROW, self.ws.max_row + 1):
            cells = {name: self.ws.cell(row, i).value for name, i in self.col.items()}
            if not cells[t.COL_ORDERER_NAME]:
                continue
            if cells[t.COL_ORDER_NO] == order_no and product in (None, cells[t.COL_PRODUCT]):
                self.ws.cell(row, self.col[column], value)
                return
        raise DemoError(f"模擬檔裡找不到單號 {order_no} {product or ''}，產生器可能改過了。")


# ---------------------------------------------------------------- the site


@dataclass(frozen=True)
class Sent:
    reply: str
    upload_id: int


class Site:
    """The site, driven as a 總部人員 in a browser would. `http` has the site as its
    base URL and does not follow redirects."""

    def __init__(self, http: httpx.Client, today: date) -> None:
        self.http = http
        self.base_url = str(http.base_url).rstrip("/")
        self.today = today

    def url(self, path: str) -> str:
        return self.base_url + path

    def log_in(self, username: str, password: str) -> None:
        resp = self.http.post("/login", data={"username": username, "password": password})
        if resp.status_code != 303:
            raise DemoError(f"用 {username} 登入失敗（HTTP {resp.status_code}），demo 帳號建好了嗎？")

    def _ok(self, resp: httpx.Response, what: str) -> httpx.Response:
        if resp.status_code != 200:
            raise DemoError(f"{what} 失敗：HTTP {resp.status_code}")
        return resp

    def has_orders(self) -> bool:
        wb = _workbook(self._ok(self.http.get("/orders.xlsx"), "下載總表"))
        return wb.active.max_row > 1

    def send(self, store_code: str, path: Path) -> Sent:
        """模擬群組上傳 as the 門市's shared account."""
        with path.open("rb") as f:
            resp = self.http.post(
                "/demo/upload",
                data={"uploader_line_id": f"Udemo-{store_code}"},
                files={"file": (path.name, f, XLSX_MEDIA_TYPE)},
            )
        page = self._ok(resp, f"上傳 {path.name}").text
        replies = _replies(page)
        upload = re.search(r'href="/uploads/(\d+)"', page)
        if len(replies) != 1 or upload is None:
            raise DemoError(f"上傳 {path.name} 之後的頁面沒有回覆")
        return Sent(replies[0], int(upload.group(1)))

    def pending_review_id(self, order_no: str) -> int:
        page = self._ok(self.http.get("/reviews"), "開待處理").text
        for review_id, row in re.findall(r'<tr class="review pending" id="review-(\d+)">(.*?)</tr>', page, re.S):
            if f">{order_no}</a>" in row:
                return int(review_id)
        raise DemoError(f"待處理頁上沒有單號 {order_no} 的待確認")

    def reject(self, review_id: int) -> str:
        resp = self.http.post("/reviews", data={"review_id": review_id, "decision": "reject"})
        if resp.status_code != 303:
            raise DemoError(f"退回待確認 #{review_id} 失敗：HTTP {resp.status_code}")
        page = self._ok(self.http.get("/reviews"), "開待處理").text
        row = re.search(rf'id="review-{review_id}">(.*?)</tr>', page, re.S)
        decision = row and re.search(r'<td class="decision">(.*?)</td>', row.group(1), re.S)
        if not decision or "已退回" not in decision.group(1):
            raise DemoError(f"待確認 #{review_id} 沒有變成已退回")
        return html.unescape(decision.group(1)).strip()

    def import_history(self, paths: tuple[Path, ...]) -> list[str]:
        files = [("files", (p.name, p.read_bytes(), XLSX_MEDIA_TYPE)) for p in paths]
        page = self._ok(self.http.post("/history/import", files=files), "匯入歷史資料").text
        return _replies(page)

    def returning_customers(self) -> dict[str, tuple]:
        """客人 with more than one 喜餅訂單, by phone: (訂單數, 門市, 訂購人姓名, 姓名不一致)."""
        ws = _workbook(self._ok(self.http.get("/customers.xlsx"), "下載客人資料"))["客人"]
        rows = ws.iter_rows(min_row=2, values_only=True)
        return {phone: (count, stores, names, odd) for phone, count, _, _, _, stores, names, odd in rows if count > 1}

    def kitchen_boxes(self) -> int:
        """全部門市 boxes to make from today for KITCHEN_DAYS days."""
        end = self.today + timedelta(days=KITCHEN_DAYS - 1)
        resp = self.http.get("/kitchen.xlsx", params={"start": self.today.isoformat(), "end": end.isoformat()})
        ws = _workbook(self._ok(resp, "下載廚房彙總")).worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        store, total = header.index("門市"), header.index("合計")
        return sum(row[total] or 0 for row in rows if row[store] == "全部門市")


def _workbook(resp: httpx.Response) -> Workbook:
    return load_workbook(io.BytesIO(resp.content))


def _replies(page: str) -> list[str]:
    return [html.unescape(m).strip() for m in re.findall(r'<div class="reply">(.*?)</div>', page, re.S)]


# ---------------------------------------------------------------- the story


class Story:
    def __init__(self, site: Site, files: DemoFiles) -> None:
        self.site = site
        self.files = files
        self.step = 0

    def title(self, text: str) -> None:
        self.step += 1
        print(f"\n== {self.step}. {text} ==")

    def send(self, store_code: str, path: Path, *, expect: tuple[str, ...] = (), forbid: tuple[str, ...] = ()) -> Sent:
        sent = self.site.send(store_code, path)
        print(f"傳檔者：{store_by_code(store_code).name} 門市帳號　檔案：{path}")
        print("LINE 回覆：")
        print(_indent(sent.reply))
        missing = [text for text in expect if text not in sent.reply]
        present = [text for text in forbid if text in sent.reply]
        if missing or present:
            raise DemoError(
                f"第 {self.step} 步的回覆跟劇本不同："
                + "".join(f"少了「{x}」；" for x in missing)
                + "".join(f"不該有「{x}」；" for x in present)
            )
        return sent

    def look(self, *paths: str) -> None:
        for path in paths:
            print(f"看這裡：{self.site.url(path)}")

    def kitchen(self, before: int | None = None) -> int:
        boxes = self.site.kitchen_boxes()
        change = "" if before is None else f"（{'沒變' if boxes == before else f'原本 {before} 盒'}）"
        print(f"廚房彙總：今天起 {KITCHEN_DAYS} 天共 {boxes} 盒{change}")
        return boxes

    def customers(self, customers: dict[str, tuple]) -> None:
        for phone, (count, stores, names, odd) in customers.items():
            print(f"・{phone}：{count} 張單，門市 {stores}，訂購人 {names}{'（姓名不一致）' if odd else ''}")

    def run(self) -> None:
        f = self.files
        self.title("門市傳了一份有錯的檔案，回覆逐列指出問題")
        sent = self.send("TC01", f.with_errors, expect=("沒有填單號", "盒數填了 0"))
        boxes = self.kitchen()
        self.look(f"/uploads/{sent.upload_id}")

        self.title("門市修正後重傳，廚房彙總跟著更新")
        self.send("TC01", f.fixed, expect=("新單 2 張",), forbid=("沒有收",))
        boxes = self.kitchen(boxes)
        self.look("/kitchen", "/orders?store=TC01")

        self.title("門市追加盒數")
        self.send("TC01", f.added, expect=(f"單號 {ADD_ORDER_NO}：追加",), forbid=("待總部確認",))
        boxes = self.kitchen(boxes)
        self.look("/kitchen")

        self.title("一筆過了修改截止點的修改，變成待確認")
        self.send("TC01", f.late_change, expect=(f"單號 {LATE_ORDER_NO}：追加", "待總部確認"))
        self.kitchen(boxes)
        self.look("/reviews", "/kitchen")

        self.title("總部退回這筆修改")
        review_id = self.site.pending_review_id(LATE_ORDER_NO)
        print(f"總部人員（{DEMO_USERNAME}）在待處理頁退回單號 {LATE_ORDER_NO}：{self.site.reject(review_id)}")
        print("LINE 回覆：沒有（不推播；結果附在這家門市下一次上傳的回覆最前面）")
        self.look("/reviews")

        self.title("門市又傳了同樣的內容，收到「已被總部退回」的提醒")
        sent = self.send("TC01", f.late_change, expect=("【總部已處理您之前的待確認】", "已被總部退回"))
        self.look(f"/uploads/{sent.upload_id}")

        self.title("門市把 Excel 改回原本的數字，提醒消失")
        sent = self.send("TC01", f.reverted, expect=("沒有新單，也沒有修改",), forbid=("退回", "待總部確認"))
        self.kitchen(boxes)
        self.look(f"/uploads/{sent.upload_id}")

        self.title("客人頁看得到回頭客")
        print("其他門市也傳了今天的檔案（回覆只列第一行）：")
        for path in f.other_stores:
            store_code = path.name.split("_", 1)[0]
            print(f"・{path.name}：{self.site.send(store_code, path).reply.splitlines()[0]}")
        before = self.site.returning_customers()
        if not before:
            raise DemoError("客人資料裡沒有回頭客")
        print("回頭客（同一支電話不只一張單）：")
        self.customers(before)
        self.look("/customers")

        self.title("匯入歷史資料")
        print("convert-legacy 轉好的範本：" + "、".join(str(p) for p in f.historical))
        for path, report in zip(f.historical, self.site.import_history(f.historical), strict=True):
            print(f"{path.name} 匯入結果：")
            print(_indent(report))
        after = self.site.returning_customers()
        found = {phone: c for phone, c in after.items() if phone not in before}
        if not found:
            raise DemoError("匯入歷史資料後沒有多出回頭客")
        print("匯入後才認得出的回頭客：")
        self.customers(found)
        self.look("/history/import", "/customers")


def _indent(text: str) -> str:
    return "\n".join(f"  {line}" for line in text.splitlines())


def replay(site: Site, files: DemoFiles) -> None:
    """Log in as the demo 總部人員 and play the story on an empty database."""
    site.log_in(DEMO_USERNAME, DEMO_PASSWORD)
    if site.has_orders():
        raise DemoError("資料庫裡已經有訂單，劇本要從空的資料庫開始。本機請加 --reset 重跑。")
    Story(site, files).run()


# ---------------------------------------------------------------- entry point


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="照順序重播 demo 劇本（9 步），印出每一步的 LINE 回覆和該看的網頁")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help=f"網站網址（預設 {DEFAULT_BASE_URL}）")
    parser.add_argument("--reset", action="store_true", help="先清空資料庫（本機用 docker compose 的 init 服務）")
    parser.add_argument("--database-url", default=None, help="--reset 時直接清空這個資料庫（部署環境要搭配 --wipe-remote）")
    parser.add_argument("--wipe-remote", action="store_true", help="確定要清空不在本機的資料庫")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT_DIR, help=f"模擬檔輸出到哪裡（預設 {DEFAULT_OUT_DIR}/）")
    args = parser.parse_args(argv)

    try:
        plan = reset_plan(args.base_url, args.database_url, wipe_remote=args.wipe_remote) if args.reset else None
        now = datetime.now(TAIPEI)
        today = today_taipei(now)
        files = prepare_files(args.out, today, now)
        print(f"模擬檔已用今天（{today}）重新產生到 {args.out}/，對照表在 {args.out}/對照表.md")
        if plan is not None:
            reset_database(plan, args.database_url)
            print("資料庫已清空，重新種入門市、品項和 demo 帳號。")
        with httpx.Client(base_url=args.base_url, timeout=90) as http:
            replay(Site(http, today), files)
    except DemoError as exc:
        print(f"\ndemo 停止：{exc}", file=sys.stderr)
        return 1
    except httpx.HTTPError as exc:
        print(f"\ndemo 停止：連不到 {args.base_url}（{exc}）。本機請先 docker compose up --build。", file=sys.stderr)
        return 1
    print(f"\ndemo 跑完了。用 {DEMO_USERNAME}／{DEMO_PASSWORD} 登入 {args.base_url} 看上面的網頁。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
