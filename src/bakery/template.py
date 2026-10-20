"""The 範本: the one .xlsx layout every 門市 fills in.

Layout (fixed so the parser can rely on it):

- A1/B1 門市代碼, C1/D1 門市名稱 (display only), E1/F1 範本版本.
- Row 2: a one-line instruction for store staff.
- Row 3: column headers. One data row = one 訂單明細, from row 4 down.

bakery.parsing imports the labels and column names below instead of repeating them.
The 品項 dropdown lists whatever names the caller passes: the live 品項 list from the
database on the web (bakery.web.routes.catalog_admin), the seed catalog for mock data.
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from bakery.xlsx import save_workbook


class StoreLike(Protocol):
    """A 門市 as the 範本 needs it: bakery.models.Store or the seed catalog's."""

    code: str
    name: str


TEMPLATE_VERSION = "v1"
SHEET_NAME = "喜餅訂單"

STORE_CODE_LABEL = "門市代碼"
STORE_NAME_LABEL = "門市"
VERSION_LABEL = "範本版本"

STORE_CODE_LABEL_CELL = "A1"
STORE_CODE_CELL = "B1"
STORE_NAME_LABEL_CELL = "C1"
STORE_NAME_CELL = "D1"
VERSION_LABEL_CELL = "E1"
VERSION_CELL = "F1"
INSTRUCTION_CELL = "A2"
HEADER_ROW = 3
FIRST_DATA_ROW = 4
# Dropdowns cover this many data rows; plenty for a store's file.
DROPDOWN_ROWS = 1000

COL_ORDER_NO = "單號"
COL_ORDERER_NAME = "訂購人姓名"
COL_PHONE = "電話"
COL_RECEIVED_ON = "接單日"
COL_PICKUP_ON = "取貨日"
COL_BANQUET_ON = "宴客日"
COL_COUPLE_NAMES = "新人姓名"
COL_TAKEN_BY = "接單人"
COL_PRODUCT = "品項"
COL_QUANTITY = "盒數"
COL_NOTE = "備註"
COL_STATUS = "狀態"

# Order-level columns: must be identical on every row of one 單號.
ORDER_COLUMNS: tuple[str, ...] = (
    COL_ORDER_NO,
    COL_ORDERER_NAME,
    COL_PHONE,
    COL_RECEIVED_ON,
    COL_PICKUP_ON,
    COL_BANQUET_ON,
    COL_COUPLE_NAMES,
    COL_TAKEN_BY,
)
LINE_COLUMNS: tuple[str, ...] = (COL_PRODUCT, COL_QUANTITY, COL_NOTE, COL_STATUS)
COLUMNS: tuple[str, ...] = ORDER_COLUMNS + LINE_COLUMNS

DATE_COLUMNS = frozenset({COL_RECEIVED_ON, COL_PICKUP_ON, COL_BANQUET_ON})
TEXT_COLUMNS = frozenset({COL_ORDER_NO, COL_PHONE})

STATUS_CANCELLED = "取消"
STATUS_OPTIONS: tuple[str, ...] = (STATUS_CANCELLED,)

DATE_NUMBER_FORMAT = "yyyy/mm/dd"

INSTRUCTION_TEXT = (
    "一列填一個品項。同一張單的每一列，從單號到接單人都要填一樣。"
    "要取消請在「狀態」選「取消」，不要刪列。要分兩次取貨，請開兩個單號。"
)

# Fixed metadata timestamp for the distributed template file.
TEMPLATE_TIMESTAMP = datetime(2026, 10, 1, 9, 0, 0)

_COLUMN_WIDTHS = {
    COL_ORDER_NO: 10,
    COL_ORDERER_NAME: 12,
    COL_PHONE: 15,
    COL_RECEIVED_ON: 12,
    COL_PICKUP_ON: 12,
    COL_BANQUET_ON: 12,
    COL_COUPLE_NAMES: 18,
    COL_TAKEN_BY: 10,
    COL_PRODUCT: 14,
    COL_QUANTITY: 8,
    COL_NOTE: 24,
    COL_STATUS: 8,
}


def _list_validation(options: Sequence[str], prompt: str) -> DataValidation:
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=True,
        showErrorMessage=False,
    )
    dv.prompt = prompt
    dv.showInputMessage = True
    return dv


def build_template(store: StoreLike, product_names: Sequence[str], headers: Sequence[str] = COLUMNS) -> Workbook:
    """Template workbook pre-filled with `store`'s code; the 品項 dropdown offers `product_names`.

    `headers` exists so mock data can reproduce stores that reorder columns
    or add their own; the dropdowns follow the 品項 / 狀態 headers wherever
    they end up.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws[STORE_CODE_LABEL_CELL] = STORE_CODE_LABEL
    ws[STORE_CODE_CELL] = store.code
    ws[STORE_NAME_LABEL_CELL] = STORE_NAME_LABEL
    ws[STORE_NAME_CELL] = store.name
    ws[VERSION_LABEL_CELL] = VERSION_LABEL
    ws[VERSION_CELL] = TEMPLATE_VERSION
    for ref in (STORE_CODE_LABEL_CELL, STORE_NAME_LABEL_CELL, VERSION_LABEL_CELL):
        ws[ref].font = Font(bold=True)

    ws[INSTRUCTION_CELL] = INSTRUCTION_TEXT
    ws[INSTRUCTION_CELL].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(
        start_row=2, start_column=1, end_row=2, end_column=max(len(headers), 1)
    )
    ws.row_dimensions[2].height = 32

    header_fill = PatternFill("solid", fgColor="FCE4D6")
    last_row = FIRST_DATA_ROW + DROPDOWN_ROWS - 1
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = _COLUMN_WIDTHS.get(header, 12)
        if header in DATE_COLUMNS:
            ws.column_dimensions[letter].number_format = DATE_NUMBER_FORMAT
        elif header in TEXT_COLUMNS:
            ws.column_dimensions[letter].number_format = "@"

        validation = None
        if header == COL_PRODUCT:
            validation = _list_validation(product_names, "請從清單選品項")
        elif header == COL_STATUS:
            validation = _list_validation(STATUS_OPTIONS, "要取消才選「取消」，平常留空")
        if validation is not None:
            validation.add(f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}")
            ws.add_data_validation(validation)

    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    return wb


def write_template(path: Path, store: StoreLike, product_names: Sequence[str]) -> Path:
    return save_workbook(build_template(store, product_names), Path(path), TEMPLATE_TIMESTAMP)
