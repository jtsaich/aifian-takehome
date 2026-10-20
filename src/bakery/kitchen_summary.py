"""廚房彙總: the boxes to make per 取貨日 × 品項, and per 門市 within each day.

Counts what is on record. 取消 lines are left out, and so are 歷史資料. A change
waiting for 總部 was not applied, so its order still holds the content from before
it: those boxes count as they are, with their cells marked 有待確認.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from sqlalchemy import ColumnElement, Select, select
from sqlalchemy.orm import Session, joinedload

from bakery.models import Change, LineStatus, Order, OrderLine, PendingReview, Product, ReviewStatus, Store
from bakery.xlsx import workbook_bytes

DEFAULT_DAYS = 14  # from today, today included
ALL_STORES = "全部門市"
PENDING = "有待確認"
WEEKDAYS = "一二三四五六日"
SUMMARY_SHEET = "廚房彙總"
PENDING_SHEET = "有待確認的訂單"
PENDING_COLUMNS = ("取貨日", "門市", "單號")


@dataclass(frozen=True)
class Boxes:
    count: int
    pending: bool  # an order counted here has a change waiting for 總部

    def __add__(self, other: Boxes) -> Boxes:
        return Boxes(self.count + other.count, self.pending or other.pending)


_NONE = Boxes(0, False)


@dataclass(frozen=True)
class SummaryRow:
    pickup_on: date
    store: Store | None  # None: every 門市 together
    boxes: dict[str, Boxes]  # by 品項 code; a 品項 nobody ordered is missing
    total: Boxes

    @property
    def store_name(self) -> str:
        return self.store.name if self.store else ALL_STORES


@dataclass(frozen=True)
class SummaryDay:
    total: SummaryRow
    stores: tuple[SummaryRow, ...]  # by 門市代碼


@dataclass(frozen=True)
class PendingOrder:
    order_id: int
    pickup_on: date
    store: Store
    order_no: str


@dataclass(frozen=True)
class KitchenSummary:
    start: date
    end: date
    products: tuple[Product, ...]  # the columns: 品項 with boxes in the range, by code
    days: tuple[SummaryDay, ...]  # only days with boxes to make
    pending_orders: tuple[PendingOrder, ...]  # in the range, by 取貨日, 門市, 單號

    @property
    def columns(self) -> tuple[str, ...]:
        return ("取貨日", "星期", "門市", *(p.name for p in self.products), "合計")


def weekday(d: date) -> str:
    return WEEKDAYS[d.weekday()]


def date_range(today: date, start: date | None, end: date | None) -> tuple[date, date]:
    """The range asked for: from `today` and DEFAULT_DAYS long unless given, in order."""
    start = start or today
    end = end or start + timedelta(days=DEFAULT_DAYS - 1)
    return (end, start) if end < start else (start, end)


def awaiting_review() -> Select[tuple[int]]:
    """Ids of orders with a 待確認 that 總部 has neither confirmed nor 退回: what 有待確認 means."""
    return (
        select(Change.order_id)
        .join(PendingReview, PendingReview.change_id == Change.id)
        .where(PendingReview.status == ReviewStatus.PENDING)
    )


def _counted(start: date, end: date) -> ColumnElement[bool]:
    """Orders the kitchen makes, picked up in the range."""
    return Order.pickup_on.between(start, end) & Order.is_historical.is_(False)


def kitchen_summary(session: Session, start: date, end: date) -> KitchenSummary:
    rows = session.execute(
        select(
            Order.pickup_on,
            Order.store_code,
            OrderLine.product_code,
            OrderLine.quantity,
            Order.id.in_(awaiting_review()),
        )
        .join(OrderLine.order)
        .where(_counted(start, end), OrderLine.status == LineStatus.ACTIVE)
    )
    by_store: dict[date, dict[str, dict[str, Boxes]]] = defaultdict(lambda: defaultdict(dict))
    for pickup_on, store_code, product_code, quantity, pending in rows:
        cells = by_store[pickup_on][store_code]
        cells[product_code] = cells.get(product_code, _NONE) + Boxes(quantity, pending)

    stores = {s.code: s for s in session.scalars(select(Store))}
    codes = {code for day in by_store.values() for cells in day.values() for code in cells}
    products = session.scalars(select(Product).where(Product.code.in_(codes)).order_by(Product.code))
    days = []
    for pickup_on in sorted(by_store):
        store_rows = tuple(
            _row(pickup_on, stores[code], cells) for code, cells in sorted(by_store[pickup_on].items())
        )
        day_cells: dict[str, Boxes] = {}
        for r in store_rows:
            for code, boxes in r.boxes.items():
                day_cells[code] = day_cells.get(code, _NONE) + boxes
        days.append(SummaryDay(_row(pickup_on, None, day_cells), store_rows))

    pending = session.scalars(
        select(Order)
        .where(_counted(start, end), Order.id.in_(awaiting_review()))
        .options(joinedload(Order.store))
        .order_by(Order.pickup_on, Order.store_code, Order.order_no)
    )
    return KitchenSummary(
        start=start,
        end=end,
        products=tuple(products),
        days=tuple(days),
        pending_orders=tuple(PendingOrder(o.id, o.pickup_on, o.store, o.order_no) for o in pending),
    )


def _row(pickup_on: date, store: Store | None, cells: dict[str, Boxes]) -> SummaryRow:
    total = _NONE
    for boxes in cells.values():
        total += boxes
    return SummaryRow(pickup_on, store, cells, total)


_DATE_FORMAT = "yyyy/mm/dd"
_PENDING_FILL = PatternFill("solid", fgColor="FFF3CD")


def kitchen_summary_xlsx(summary: KitchenSummary, now: datetime) -> bytes:
    """The 廚房彙總 page as a workbook. Each day's 門市 rows are grouped under its
    全部門市 row and folded, as on the page; 有待確認 is a comment on the number."""
    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY_SHEET
    ws.append(summary.columns)
    ws.sheet_properties.outlinePr.summaryBelow = False
    for day in summary.days:
        for row in (day.total, *day.stores):
            cells = [*(row.boxes.get(p.code) for p in summary.products), row.total]
            ws.append(
                (
                    row.pickup_on,
                    weekday(row.pickup_on),
                    row.store_name,
                    *(boxes.count if boxes else None for boxes in cells),
                )
            )
            r = ws.max_row
            ws.cell(r, 1).number_format = _DATE_FORMAT
            for col, boxes in enumerate(cells, start=4):
                if boxes is not None and boxes.pending:
                    cell = ws.cell(r, col)
                    cell.comment = Comment(PENDING, "bakery-orders")
                    cell.fill = _PENDING_FILL
            if row.store is None:
                for cell in ws[r]:
                    cell.font = Font(bold=True)
            else:
                ws.row_dimensions[r].outlineLevel = 1
                ws.row_dimensions[r].hidden = True
    ws.freeze_panes = "D2"

    pending = wb.create_sheet(PENDING_SHEET)
    pending.append(PENDING_COLUMNS)
    for o in summary.pending_orders:
        pending.append((o.pickup_on, o.store.name, o.order_no))
        pending.cell(pending.max_row, 1).number_format = _DATE_FORMAT
    pending.freeze_panes = "A2"
    return workbook_bytes(wb, now.replace(tzinfo=None))
