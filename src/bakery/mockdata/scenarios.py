"""The demo's mock files: what each store uploads on day 1 and day 2, plus legacy files.

Seeded orders use fixed 品項 and 盒數 so the 對照表 reads the same on every
run; names, phones and filler orders come from the seeded RNG. All dates are
offsets from `today` (day 1).
"""

from __future__ import annotations

import copy
import random
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from bakery import template as t
from bakery.catalog import STORES, Store, product_names, store_by_code
from bakery.mockdata.expectations import (
    ChangeExpectation,
    ChangeKind,
    FieldDiff,
    FileExpectation,
    Finding,
    FindingKind,
    LegacyExpectation,
    LegacyOrder,
    Manifest,
    Outcome,
    ReturningCustomer,
    Seeded,
    Severity,
    Tolerance,
    ToleranceKind,
    render_markdown,
)
from bakery.mockdata.people import People, format_phone
from bakery.mockdata.sheet import (
    DATE_ROC,
    MockLine,
    MockOrder,
    roc_text,
    slash_text,
    write_orders,
)
from bakery.xlsx import save_workbook

ASSUMPTIONS: dict[str, int] = {
    "cutoff_days": 7,
    "pickup_too_soon_days": 3,
    "quantity_warning_over": 300,
}

# Filler orders stay clear of every warning threshold above.
_FILLER_PICKUP_DAYS = (12, 60)
_FILLER_BOXES = (20, 200)

# Offset (days after day 1) of the TC01 orders whose day-2 changes fall past the cutoff.
PAST_CUTOFF_PICKUP_OFFSET = 6

MEMO_SHEET_NAME = "備忘"
EXTRA_COLUMN_NAME = "訂金"
PHONE_SYNONYM = "手機"
UNKNOWN_PRODUCT_TEXT = "月下"
UNPARSEABLE_DATE_TEXT = "下週六"


@dataclass
class _StoreStyle:
    prefix: str
    start: int
    phone_style: str
    staff: tuple[str, ...]


_STYLES: dict[str, _StoreStyle] = {
    "TC01": _StoreStyle("", 101, "dash", ("美玲", "佩君", "志豪")),
    "TC02": _StoreStyle("B", 231, "plain", ("淑惠", "雅琪")),
    "TC03": _StoreStyle("", 1501, "space", ("秀蘭", "宛臻", "建志")),
    "CH01": _StoreStyle("CH-", 88, "dash", ("麗華", "怡伶")),
    "CH02": _StoreStyle("", 457, "plain", ("素珍", "家瑋", "玉婷")),
    "NT01": _StoreStyle("", 12, "dash", ("阿月", "惠君")),
}


@dataclass
class _Ctx:
    """Generation state for one store."""

    store: Store
    style: _StoreStyle
    rng: random.Random
    people: People
    today: date
    _next: int = 0

    def __post_init__(self) -> None:
        self._next = self.style.start

    def next_no(self) -> str:
        no = f"{self.style.prefix}{self._next:04d}"
        self._next += 1
        return no

    def day(self, offset: int) -> date:
        return self.today + timedelta(days=offset)

    def phone(self, digits: str | None = None, style: str | None = None) -> str:
        return format_phone(digits or self.people.mobile_digits(), style or self.style.phone_style)

    def order(
        self,
        lines: list[tuple[str, int]],
        pickup_in: int,
        *,
        order_no: str | None = "auto",
        banquet_after: int | None = 1,
        phone: str | None = None,
        orderer: str | None = None,
    ) -> MockOrder:
        """An order with fixed 品項/盒數 and fake people."""
        pickup = self.day(pickup_in)
        return MockOrder(
            order_no=self.next_no() if order_no == "auto" else order_no,
            orderer_name=orderer or self.people.orderer(),
            phone=phone or self.phone(),
            received_on=self.day(-self.rng.randint(0, 6)),
            pickup_on=pickup,
            banquet_on=None if banquet_after is None else pickup + timedelta(days=banquet_after),
            couple_names=self.people.couple(),
            taken_by=self.rng.choice(self.style.staff),
            lines=[MockLine(p, q) for p, q in lines],
        )

    def filler(self, pickup_in: int | None = None) -> MockOrder:
        """A random, clean order."""
        n = self.rng.choice((1, 1, 2, 2, 3))
        products = self.rng.sample(product_names(), n)
        lo, hi = _FILLER_BOXES
        lines = [(p, self.rng.randrange(lo // 10, hi // 10 + 1) * 10) for p in products]
        order = self.order(
            lines,
            pickup_in if pickup_in is not None else self.rng.randint(*_FILLER_PICKUP_DAYS),
            banquet_after=self.rng.choice((0, 1, 1, 2)) if self.rng.random() < 0.7 else None,
        )
        order.lines[0].note = self.people.note()
        return order


# ---------------------------------------------------------------- helpers


def _at(order: MockOrder) -> str:
    """Location prefix for reply text, e.g. 「喜餅訂單 第 5、6 列（單號 0458）」."""
    rows = "、".join(str(r) for r in order.rows)
    no = f"（單號 {order.order_no}）" if order.order_no else ""
    return f"{order.sheet} 第 {rows} 列{no}"


def _summary(store: Store, orders: int, lines: int, rejected_rows: int, warnings: int) -> str:
    text = f"{store.name}（{store.code}）：收到 {orders} 張喜餅訂單、{lines} 筆訂單明細。"
    if rejected_rows:
        text += f"有 {rejected_rows} 列沒有收，請改好後再傳一次。"
    if warnings:
        text += f"有 {warnings} 個地方已先收下，請再確認。"
    return text


def _new_sheet(store: Store, headers: list[str] | None = None) -> Workbook:
    return t.build_template(store, headers or list(t.COLUMNS))


def _expectation(
    ctx: _Ctx,
    relpath: str,
    orders: list[MockOrder],
    *,
    seeded: list[Seeded],
    findings: list[Finding],
    tolerances: list[Tolerance],
    rejected: list[MockOrder],
) -> FileExpectation:
    """Day-1 expectation: everything not in `rejected` is a new, accepted order."""
    rejected_ids = {id(o) for o in rejected}
    accepted = [o for o in orders if id(o) not in rejected_ids]
    warnings = sum(1 for f in findings if f.severity == Severity.WARNING)
    rejected_rows = sorted(r for o in rejected for r in o.rows)
    return FileExpectation(
        file=relpath,
        stage="day1",
        store_code=ctx.store.code,
        upload_on=ctx.today.isoformat(),
        process_after=[],
        seeded=seeded,
        accepted_orders=[o.order_no for o in accepted],
        accepted_lines=sum(len(o.lines) for o in accepted),
        rejected_orders=[o.order_no for o in rejected if o.order_no],
        rejected_rows=rejected_rows,
        findings=findings,
        tolerances=tolerances,
        reply_summary=_summary(
            ctx.store, len(accepted), sum(len(o.lines) for o in accepted), len(rejected_rows), warnings
        ),
        new_orders=[o.order_no for o in accepted],
    )


def _roc_tolerances(orders: list[MockOrder]) -> list[Tolerance]:
    return [
        Tolerance(
            ToleranceKind.ROC_DATE,
            o.sheet,
            o.rows,
            {
                "接單日寫法": roc_text(o.received_on),
                "接單日應解析為": o.received_on.isoformat(),
                "取貨日寫法": roc_text(o.pickup_on),
                "取貨日應解析為": o.pickup_on.isoformat(),
            },
        )
        for o in orders
    ]


@dataclass
class _Day1:
    workbook: Workbook
    expectation: FileExpectation
    orders: dict[str, MockOrder] = field(default_factory=dict)  # by 單號, as written


# ---------------------------------------------------------------- day 1


def _tc01(ctx: _Ctx, relpath: str) -> _Day1:
    """漏單號 + 盒數 0. Also the orders that day 2 modifies."""
    mom_phone = ctx.phone()
    mom = ctx.people.female()
    o101 = ctx.order([("龍鳳喜餅", 80), ("鳳梨酥禮盒", 40)], 20, phone=mom_phone, orderer=mom)
    o102 = ctx.order([("鳳梨酥禮盒", 60)], 18)
    o103 = ctx.order([("綠豆椪禮盒", 50), ("太陽餅禮盒", 30)], 20, banquet_after=None)
    o104 = ctx.order([("綜合西式喜餅", 100), ("蛋黃酥禮盒", 50)], 30)
    # Same phone and orderer as 0101: a mother ordering for her second daughter.
    o105 = ctx.order([("咖哩大餅", 70)], 22, phone=mom_phone, orderer=mom)
    missing = ctx.order([("太陽餅禮盒", 60)], 26, order_no=None)
    o106 = ctx.order([("綠豆椪禮盒", 50), ("龍鳳喜餅", 30)], PAST_CUTOFF_PICKUP_OFFSET)
    o107 = ctx.order([("太陽餅禮盒", 40)], PAST_CUTOFF_PICKUP_OFFSET, banquet_after=1)
    o108 = ctx.order([("蛋黃酥禮盒", 60)], 25)
    o109 = ctx.order([("龍鳳喜餅", 100)], 25, banquet_after=None)
    o110 = ctx.filler()
    o111 = ctx.filler()
    o112 = ctx.order([("太陽餅禮盒", 0)], 28)
    orders = [o101, o102, o103, o104, o105, missing, o106, o107, o108, o109, o110, o111, o112]

    wb = _new_sheet(ctx.store)
    write_orders(wb.active, orders, t.COLUMNS)
    findings = [
        Finding(
            FindingKind.MISSING_ORDER_NO,
            Severity.ERROR,
            missing.sheet,
            missing.rows,
            [],
            f"{_at(missing)}：沒有填單號，這一列沒有收。請補上單號再傳一次。",
        ),
        Finding(
            FindingKind.INVALID_QUANTITY,
            Severity.ERROR,
            o112.sheet,
            o112.rows,
            [o112.order_no],
            f"{_at(o112)}：盒數填了 0，要填 1 以上的整數，這一列沒有收。",
        ),
    ]
    exp = _expectation(
        ctx,
        relpath,
        orders,
        seeded=[Seeded.MISSING_ORDER_NO, Seeded.INVALID_QUANTITY],
        findings=findings,
        tolerances=[],
        rejected=[missing, o112],
    )
    return _Day1(wb, exp, {o.order_no: o for o in orders if o.order_no})


def _tc02(ctx: _Ctx, relpath: str) -> _Day1:
    """民國年日期 + 取貨日離今天太近."""
    o1 = ctx.order([("龍鳳喜餅", 60), ("綠豆椪禮盒", 40)], 24)
    o2 = ctx.order([("太陽餅禮盒", 80)], 31)
    o3 = ctx.filler()
    o4 = ctx.order([("鳳梨酥禮盒", 50)], 2)
    o5 = ctx.filler()
    for o in (o1, o2, o4):
        o.date_style = DATE_ROC
    orders = [o1, o2, o3, o4, o5]
    wb = _new_sheet(ctx.store)
    write_orders(wb.active, orders, t.COLUMNS)
    findings = [
        Finding(
            FindingKind.PICKUP_TOO_SOON,
            Severity.WARNING,
            o4.sheet,
            o4.rows,
            [o4.order_no],
            f"{_at(o4)}：取貨日 {slash_text(o4.pickup_on)} 只剩 2 天，已先收下，請確認中央廚房來得及做。",
        )
    ]
    exp = _expectation(
        ctx,
        relpath,
        orders,
        seeded=[Seeded.ROC_DATE, Seeded.PICKUP_TOO_SOON],
        findings=findings,
        tolerances=_roc_tolerances([o1, o2, o4]),
        rejected=[],
    )
    return _Day1(wb, exp, {o.order_no: o for o in orders})


_TC03_HEADERS = [
    t.COL_ORDER_NO,
    t.COL_PRODUCT,
    t.COL_QUANTITY,
    t.COL_ORDERER_NAME,
    PHONE_SYNONYM,
    t.COL_RECEIVED_ON,
    t.COL_PICKUP_ON,
    t.COL_BANQUET_ON,
    t.COL_COUPLE_NAMES,
    t.COL_TAKEN_BY,
    t.COL_NOTE,
    t.COL_STATUS,
]


def _tc03(ctx: _Ctx, relpath: str) -> _Day1:
    """品項用別名 + 不在清單的品項; columns reordered and 電話 headed 手機."""
    o1 = ctx.order([("龍鳳餅", 120), ("綠豆椪禮盒", 30)], 19)
    o2 = ctx.order([("鳳梨酥", 90)], 27)
    o3 = ctx.order([("西餅", 60), ("太陽餅禮盒", 40)], 33)
    o4 = ctx.order([(UNKNOWN_PRODUCT_TEXT, 50)], 40)
    o5 = ctx.filler()
    orders = [o1, o2, o3, o4, o5]
    wb = _new_sheet(ctx.store, _TC03_HEADERS)
    columns = [t.COL_PHONE if h == PHONE_SYNONYM else h for h in _TC03_HEADERS]
    write_orders(wb.active, orders, columns)
    sheet = wb.active.title
    alias_rows = [(o1.lines[0], "龍鳳喜餅"), (o2.lines[0], "鳳梨酥禮盒"), (o3.lines[0], "綜合西式喜餅")]
    tolerances = [
        Tolerance(ToleranceKind.PRODUCT_ALIAS, sheet, [line.row], {"寫法": line.product, "應解析為": canonical})
        for line, canonical in alias_rows
    ]
    tolerances += [
        Tolerance(
            ToleranceKind.REORDERED_COLUMNS,
            sheet,
            [],
            {"欄位順序": "、".join(_TC03_HEADERS)},
        ),
        Tolerance(ToleranceKind.HEADER_SYNONYM, sheet, [], {"欄名": PHONE_SYNONYM, "應對應到": t.COL_PHONE}),
    ]
    findings = [
        Finding(
            FindingKind.UNKNOWN_PRODUCT,
            Severity.ERROR,
            o4.sheet,
            o4.rows,
            [o4.order_no],
            f"{_at(o4)}：品項「{UNKNOWN_PRODUCT_TEXT}」不在商品清單裡，這一列沒有收。"
            "請改成清單上的品項；如果是新品項，請總部加進清單後再傳一次。",
        )
    ]
    exp = _expectation(
        ctx,
        relpath,
        orders,
        seeded=[Seeded.PRODUCT_ALIAS, Seeded.UNKNOWN_PRODUCT, Seeded.REORDERED_COLUMNS, Seeded.HEADER_SYNONYM],
        findings=findings,
        tolerances=tolerances,
        rejected=[o4],
    )
    return _Day1(wb, exp, {o.order_no: o for o in orders})


def _ch01(ctx: _Ctx, relpath: str) -> _Day1:
    """同單號訂單層欄位不一致 + 日期看不懂."""
    o1 = ctx.filler()
    o2 = ctx.order([("龍鳳喜餅", 100), ("鳳梨酥禮盒", 50)], 23)
    second_pickup = o2.pickup_on + timedelta(days=7)
    o2.lines[1].cells[t.COL_PICKUP_ON] = second_pickup
    o3 = ctx.order([("咖哩大餅", 40)], 29)
    o3.cells[t.COL_PICKUP_ON] = UNPARSEABLE_DATE_TEXT
    o4 = ctx.filler()
    o5 = ctx.filler()
    orders = [o1, o2, o3, o4, o5]
    wb = _new_sheet(ctx.store)
    write_orders(wb.active, orders, t.COLUMNS)
    findings = [
        Finding(
            FindingKind.INCONSISTENT_ORDER_FIELDS,
            Severity.ERROR,
            o2.sheet,
            o2.rows,
            [o2.order_no],
            f"{_at(o2)}：同一張單的取貨日不一樣（{slash_text(o2.pickup_on)}、{slash_text(second_pickup)}），"
            "整張單沒有收。同一張單的每一列，從單號到接單人都要填一樣；要分兩次取貨，請開兩個單號。",
        ),
        Finding(
            FindingKind.UNPARSEABLE_DATE,
            Severity.ERROR,
            o3.sheet,
            o3.rows,
            [o3.order_no],
            f"{_at(o3)}：取貨日寫「{UNPARSEABLE_DATE_TEXT}」，系統看不懂，這一列沒有收。"
            "請填像 2026/11/08 這樣的日期。",
        ),
    ]
    exp = _expectation(
        ctx,
        relpath,
        orders,
        seeded=[Seeded.INCONSISTENT_ORDER_FIELDS, Seeded.UNPARSEABLE_DATE],
        findings=findings,
        tolerances=[],
        rejected=[o2, o3],
    )
    return _Day1(wb, exp, {o.order_no: o for o in orders})


_CH02_HEADERS = [
    t.COL_ORDER_NO,
    t.COL_ORDERER_NAME,
    t.COL_PHONE,
    EXTRA_COLUMN_NAME,
    *t.COLUMNS[3:],
]


def _ch02(ctx: _Ctx, relpath: str) -> _Day1:
    """取貨日已經過了 + 自己加的「訂金」欄, plus the remaining warnings."""
    o1 = ctx.filler()
    o2 = ctx.order([("綠豆椪禮盒", 60), ("蛋黃酥禮盒", 40)], -3)
    o2.received_on = ctx.day(-12)
    o3 = ctx.order([("太陽餅禮盒", 70)], 21)
    # Same order typed twice under two 單號.
    o4 = copy.deepcopy(o3)
    o4.order_no = ctx.next_no()
    bad_digits = ctx.people.mobile_digits()[:8]
    o5 = ctx.order([("咖哩大餅", 50)], 35, phone=f"{bad_digits[:4]}-{bad_digits[4:]}")
    o6 = ctx.order([("龍鳳喜餅", 400)], 44)
    orders = [o1, o2, o3, o4, o5, o6]
    for o, deposit in zip(orders, (3000, 2000, None, None, 1500, 10000), strict=True):
        if deposit:
            o.cells[EXTRA_COLUMN_NAME] = deposit
    wb = _new_sheet(ctx.store, _CH02_HEADERS)
    write_orders(wb.active, orders, _CH02_HEADERS)
    findings = [
        Finding(
            FindingKind.PICKUP_IN_PAST,
            Severity.WARNING,
            o2.sheet,
            o2.rows,
            [o2.order_no],
            f"{_at(o2)}：取貨日 {slash_text(o2.pickup_on)} 已經過了，已先收下，請確認日期有沒有打錯。",
        ),
        Finding(
            FindingKind.DUPLICATE_PHONE_PICKUP,
            Severity.WARNING,
            o3.sheet,
            o3.rows + o4.rows,
            [o3.order_no, o4.order_no],
            f"{o3.sheet} 第 {o3.rows[0]}、{o4.rows[0]} 列（單號 {o3.order_no}、{o4.order_no}）："
            "同一支電話、同一天取貨，已先收下，請確認是不是重複登記。",
        ),
        Finding(
            FindingKind.PHONE_FORMAT,
            Severity.WARNING,
            o5.sheet,
            o5.rows,
            [o5.order_no],
            f"{_at(o5)}：電話「{o5.phone}」好像少了幾碼，已先收下，請確認號碼。",
        ),
        Finding(
            FindingKind.QUANTITY_OVER_THRESHOLD,
            Severity.WARNING,
            o6.sheet,
            o6.rows,
            [o6.order_no],
            f"{_at(o6)}：龍鳳喜餅 400 盒比平常多很多，已先收下，請確認有沒有多打一個 0。",
        ),
    ]
    tolerances = [
        Tolerance(
            ToleranceKind.EXTRA_COLUMN,
            wb.active.title,
            [],
            {"欄名": EXTRA_COLUMN_NAME, "處理": "這一欄不讀，其他欄照常收"},
        )
    ]
    exp = _expectation(
        ctx,
        relpath,
        orders,
        seeded=[
            Seeded.EXTRA_COLUMN,
            Seeded.PICKUP_IN_PAST,
            Seeded.DUPLICATE_PHONE_PICKUP,
            Seeded.PHONE_FORMAT,
            Seeded.QUANTITY_OVER_THRESHOLD,
        ],
        findings=findings,
        tolerances=tolerances,
        rejected=[],
    )
    return _Day1(wb, exp, {o.order_no: o for o in orders})


_MEMO_LINES = (
    "週一公休，週日下午人手少",
    "龍鳳喜餅的提袋快用完了，記得跟總部叫貨",
    "有客人問喜餅可不可以宅配：目前不行",
    "月底盤點",
)


def _nt01(ctx: _Ctx, relpath: str, cross_store_digits: str, cross_store_name: str, avoid_pickup: date) -> _Day1:
    """多了「備忘」分頁; phones typed in several formats."""
    pickup_in = (avoid_pickup - ctx.today).days + 9
    o1 = ctx.order(
        [("綜合西式喜餅", 80)],
        pickup_in,
        phone=ctx.phone(cross_store_digits, "intl"),
        orderer=f"{cross_store_name[0]}小姐",
    )
    o2 = ctx.filler()
    o2.phone = ctx.phone(ctx.people.landline_digits("049"), "dash")
    o3 = ctx.filler()
    o4 = ctx.filler()
    o4.phone = ctx.phone(style="space")
    o5 = ctx.filler()
    orders = [o1, o2, o3, o4, o5]
    wb = _new_sheet(ctx.store)
    write_orders(wb.active, orders, t.COLUMNS)
    memo = wb.create_sheet(MEMO_SHEET_NAME)
    memo["A1"] = "店內備忘"
    memo["A1"].font = Font(bold=True)
    for i, text in enumerate(_MEMO_LINES, start=2):
        memo.cell(row=i, column=1, value=text)
    findings = [
        Finding(
            FindingKind.SKIPPED_SHEET,
            Severity.INFO,
            MEMO_SHEET_NAME,
            [],
            [],
            f"略過分頁：{MEMO_SHEET_NAME}（這個分頁的欄位跟範本不一樣，系統沒有讀）。",
        )
    ]
    exp = _expectation(
        ctx,
        relpath,
        orders,
        seeded=[Seeded.MEMO_SHEET],
        findings=findings,
        tolerances=[],
        rejected=[],
    )
    return _Day1(wb, exp, {o.order_no: o for o in orders})


# ---------------------------------------------------------------- day 2


def _tc01_day2(ctx: _Ctx, day1: _Day1, day1_path: str, relpath: str) -> tuple[Workbook, FileExpectation]:
    """TC01 sends its running file again with changes; see the 對照表 for each one."""
    d1 = day1.orders
    o = {no: copy.deepcopy(order) for no, order in d1.items()}
    day2 = ctx.day(1)
    changes: list[ChangeExpectation] = []

    def change(no, kinds, diffs, reply, outcome=Outcome.APPLIED):
        past = (o[no].pickup_on - day2).days < ASSUMPTIONS["cutoff_days"] or (
            d1[no].pickup_on - day2
        ).days < ASSUMPTIONS["cutoff_days"]
        changes.append(ChangeExpectation(no, kinds, diffs, past, outcome, reply))

    no = "0101"
    o[no].lines[0].quantity = 120
    change(no, [ChangeKind.ADD_QTY], [FieldDiff(t.COL_QUANTITY, "80", "120", "龍鳳喜餅")], f"單號 {no}：追加，龍鳳喜餅 80→120 盒")

    no = "0102"
    o[no].lines[0].quantity = 40
    change(no, [ChangeKind.REDUCE_QTY], [FieldDiff(t.COL_QUANTITY, "60", "40", "鳳梨酥禮盒")], f"單號 {no}：減量，鳳梨酥禮盒 60→40 盒")

    no = "0103"
    old, new = d1[no].pickup_on, d1[no].pickup_on + timedelta(days=7)
    o[no].pickup_on = new
    change(
        no,
        [ChangeKind.PICKUP_DATE],
        [FieldDiff(t.COL_PICKUP_ON, old.isoformat(), new.isoformat())],
        f"單號 {no}：改取貨日，{slash_text(old)}→{slash_text(new)}",
    )

    no = "0104"
    for line in o[no].lines:
        line.status = t.STATUS_CANCELLED
    change(
        no,
        [ChangeKind.CANCEL],
        [FieldDiff(t.COL_STATUS, "", t.STATUS_CANCELLED, line.product) for line in o[no].lines],
        f"單號 {no}：取消，綜合西式喜餅 100 盒、蛋黃酥禮盒 50 盒不再生產",
    )

    no = "0105"
    new_couple = ctx.people.couple()
    o[no].couple_names = new_couple
    change(
        no,
        [ChangeKind.INFO_FIX],
        [FieldDiff(t.COL_COUPLE_NAMES, d1[no].couple_names, new_couple)],
        f"單號 {no}：資料更正，新人姓名 {d1[no].couple_names}→{new_couple}",
    )

    no = "0106"
    o[no].lines[0].quantity = 70
    change(
        no,
        [ChangeKind.ADD_QTY],
        [FieldDiff(t.COL_QUANTITY, "50", "70", "綠豆椪禮盒")],
        f"單號 {no}：追加，綠豆椪禮盒 50→70 盒（已過截止點，待總部確認）",
        Outcome.PENDING,
    )

    # Phone changes but 品項 and 取貨日 stay: still 資料更正, not 疑似單號重複使用.
    no = "0107"
    new_banquet = d1[no].banquet_on + timedelta(days=1)
    o[no].banquet_on = new_banquet
    o[no].phone = ctx.phone()
    change(
        no,
        [ChangeKind.INFO_FIX],
        [
            FieldDiff(t.COL_PHONE, d1[no].phone, o[no].phone),
            FieldDiff(t.COL_BANQUET_ON, d1[no].banquet_on.isoformat(), new_banquet.isoformat()),
        ],
        f"單號 {no}：資料更正，電話已更新；宴客日 {slash_text(d1[no].banquet_on)}→{slash_text(new_banquet)}",
    )

    no = "0108"
    reused = ctx.order([("鳳梨酥禮盒", 80)], 40, order_no=no)
    o[no] = reused
    change(
        no,
        [ChangeKind.SUSPECTED_REUSE],
        [
            FieldDiff(t.COL_PHONE, d1[no].phone, reused.phone),
            FieldDiff(t.COL_ORDERER_NAME, d1[no].orderer_name, reused.orderer_name),
            FieldDiff(t.COL_PICKUP_ON, d1[no].pickup_on.isoformat(), reused.pickup_on.isoformat()),
            FieldDiff(t.COL_QUANTITY, "0", "80", "鳳梨酥禮盒"),
        ],
        f"單號 {no}：電話跟之前不同，取貨日和品項也變了，可能是兩張不同的單用了同一個單號。"
        "這張單先照舊，待總部確認；請檢查單號有沒有重複。",
        Outcome.PENDING,
    )

    no = "0109"
    old = d1[no].pickup_on
    new = old + timedelta(days=7)
    o[no].pickup_on = new
    o[no].lines.append(MockLine("鳳梨酥禮盒", 30))
    change(
        no,
        [ChangeKind.ADD_QTY, ChangeKind.PICKUP_DATE],
        [FieldDiff(t.COL_QUANTITY, "0", "30", "鳳梨酥禮盒"), FieldDiff(t.COL_PICKUP_ON, old.isoformat(), new.isoformat())],
        f"單號 {no}：追加＋改取貨日，新增鳳梨酥禮盒 30 盒；取貨日 {slash_text(old)}→{slash_text(new)}",
    )

    unchanged = ["0110"]
    absent = ["0111"]
    new_order = ctx.filler()
    new_order.received_on = day2
    file_orders = [o[n] for n in ("0101", "0102", "0103", "0104", "0105", "0106", "0107", "0108", "0109", "0110")]
    file_orders.append(new_order)

    wb = _new_sheet(ctx.store)
    write_orders(wb.active, file_orders, t.COLUMNS)
    pending = sum(1 for c in changes if c.outcome == Outcome.PENDING)
    exp = FileExpectation(
        file=relpath,
        stage="day2",
        store_code=ctx.store.code,
        upload_on=day2.isoformat(),
        process_after=[day1_path],
        seeded=[],
        accepted_orders=[x.order_no for x in file_orders],
        accepted_lines=sum(len(x.lines) for x in file_orders),
        rejected_orders=[],
        rejected_rows=[],
        findings=[],
        tolerances=[],
        reply_summary=(
            f"{ctx.store.name}（{ctx.store.code}）：新單 1 張；修改 {len(changes)} 張，"
            f"其中 {pending} 張待總部確認。"
        ),
        new_orders=[new_order.order_no],
        changes=changes,
        unchanged_orders=unchanged,
        absent_orders=absent,
    )
    return wb, exp


def _resend_expectation(day1: FileExpectation, store: Store, relpath: str, day2: date) -> FileExpectation:
    return FileExpectation(
        file=relpath,
        stage="day2",
        store_code=store.code,
        upload_on=day2.isoformat(),
        process_after=[day1.file],
        seeded=[],
        accepted_orders=list(day1.accepted_orders),
        accepted_lines=day1.accepted_lines,
        rejected_orders=[],
        rejected_rows=[],
        findings=[f for f in day1.findings if f.kind == FindingKind.SKIPPED_SHEET],
        tolerances=[],
        reply_summary=f"{store.name}（{store.code}）：這份檔案的內容之前都收過了，沒有新單，也沒有修改。",
        new_orders=[],
        unchanged_orders=list(day1.accepted_orders),
    )


# ---------------------------------------------------------------- legacy

_LEGACY_TC02_ITEM_COLUMNS = {"龍鳳": "龍鳳喜餅", "鳳梨酥": "鳳梨酥禮盒", "綠豆椪": "綠豆椪禮盒"}
_LEGACY_OTHER = {"太陽餅禮盒": "太陽餅", "蛋黃酥禮盒": "蛋黃酥", "咖哩大餅": "咖哩餅"}


def _legacy_dates(ctx: _Ctx, offsets: tuple[int, ...]) -> list[tuple[date, date]]:
    out = []
    for off in offsets:
        received = ctx.day(-off - ctx.rng.randint(0, 4))
        out.append((received, received + timedelta(days=ctx.rng.randint(14, 30))))
    return out


def _legacy_tc02(ctx: _Ctx, relpath: str, returning_digits: str) -> tuple[Workbook, LegacyExpectation]:
    """One row per order; the main 品項 each get a column, the rest go in 其他 as text."""
    wb = Workbook()
    ws = wb.active
    ws.title = "喜餅登記"
    headers = ["接單日期", "客人", "電話", *_LEGACY_TC02_ITEM_COLUMNS, "其他", "取貨日", "備註"]
    ws.append(headers)
    plan = [
        {"龍鳳": 60, "鳳梨酥": 30},
        {"綠豆椪": 80},
        {"龍鳳": 50, "其他": ("太陽餅禮盒", 20)},
        {"鳳梨酥": 100},
        {"龍鳳": 40, "綠豆椪": 40, "其他": ("蛋黃酥禮盒", 30)},
        {"其他": ("咖哩大餅", 60)},
        {"龍鳳": 120},
    ]
    dates = _legacy_dates(ctx, (150, 130, 115, 100, 80, 62, 45))
    orders: list[LegacyOrder] = []
    for i, (items, (received, pickup)) in enumerate(zip(plan, dates, strict=True)):
        name = ctx.people.orderer()
        digits = returning_digits if i == 4 else ctx.people.mobile_digits()
        phone = format_phone(digits, "space" if i == 4 else "plain")
        row = [roc_text(received), name, phone]
        lines: dict[str, int] = {}
        for col, canonical in _LEGACY_TC02_ITEM_COLUMNS.items():
            qty = items.get(col)
            row.append(qty)
            if qty:
                lines[canonical] = qty
        other = items.get("其他")
        row.append(f"{_LEGACY_OTHER[other[0]]}{other[1]}盒" if other else None)
        if other:
            lines[other[0]] = other[1]
        row += [roc_text(pickup), ctx.people.note() or None]
        ws.append(row)
        ws.cell(row=ws.max_row, column=3).number_format = "@"
        orders.append(
            LegacyOrder(ws.title, [ws.max_row], name, phone, received.isoformat(), pickup.isoformat(), lines)
        )
    exp = LegacyExpectation(
        file=relpath,
        store_code=ctx.store.code,
        layout="一列一張單；龍鳳、鳳梨酥、綠豆椪各佔一欄填盒數，其他品項寫在「其他」欄（例如「太陽餅20盒」）；日期是民國年文字；沒有單號、宴客日、新人、接單人欄。",
        orders=orders,
        notes=[
            "沒有單號：轉換工具要替每張單補一個一看就知道是歷史資料的單號。",
            "欄名對照：龍鳳→龍鳳喜餅、鳳梨酥→鳳梨酥禮盒、綠豆椪→綠豆椪禮盒；「其他」欄的太陽餅、蛋黃酥、咖哩餅是品項別名或簡稱。",
            "取貨日都已經過了，匯入後不會出現在廚房彙總。",
        ],
    )
    return wb, exp


_LEGACY_CH01_HEADERS = ["訂購日", "姓名", "聯絡電話", "品名", "數量", "交貨日", "結婚日", "新人"]
_LEGACY_CH01_PLAN = (
    (("龍鳳", 80), ("鳳梨酥", 40)),
    (("西式喜餅", 100),),
    (("綠豆椪禮盒", 60), ("太陽餅", 30), ("蛋黃酥禮盒", 20)),
    (("龍鳳喜餅", 50),),
    (("咖哩大餅", 40), ("西餅", 40)),
    (("鳳梨酥禮盒", 120),),
)
_LEGACY_CH01_CANONICAL = {
    "龍鳳": "龍鳳喜餅",
    "鳳梨酥": "鳳梨酥禮盒",
    "西式喜餅": "綜合西式喜餅",
    "西餅": "綜合西式喜餅",
    "太陽餅": "太陽餅禮盒",
}


def _legacy_ch01(ctx: _Ctx, relpath: str, returning_digits: str) -> tuple[Workbook, LegacyExpectation]:
    """One row per item, no 單號, one sheet per month of 訂購日."""
    wb = Workbook()
    wb.remove(wb.active)
    dates = _legacy_dates(ctx, (140, 120, 95, 75, 55, 40))
    orders: list[LegacyOrder] = []
    for i, (items, (received, pickup)) in enumerate(zip(_LEGACY_CH01_PLAN, dates, strict=True)):
        sheet_name = f"{received.year - 1911}年{received.month}月"
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(sheet_name)
            ws.append(_LEGACY_CH01_HEADERS)
        name = ctx.people.orderer()
        digits = returning_digits if i == 3 else ctx.people.mobile_digits()
        phone = format_phone(digits, "plain" if i == 3 else "dash")
        banquet = pickup + timedelta(days=1)
        couple = ctx.people.couple()
        rows: list[int] = []
        lines: dict[str, int] = {}
        for product, qty in items:
            ws.append([received, name, phone, product, qty, pickup, banquet, couple])
            r = ws.max_row
            for col in (1, 6, 7):
                ws.cell(row=r, column=col).number_format = t.DATE_NUMBER_FORMAT
            ws.cell(row=r, column=3).number_format = "@"
            rows.append(r)
            lines[_LEGACY_CH01_CANONICAL.get(product, product)] = qty
        orders.append(LegacyOrder(ws.title, rows, name, phone, received.isoformat(), pickup.isoformat(), lines))
    exp = LegacyExpectation(
        file=relpath,
        store_code=ctx.store.code,
        layout="一列一個品項，同一張單的品項排在相鄰幾列；每個月一個分頁（例如「115年8月」）；沒有單號欄；電話欄叫「聯絡電話」，取貨日叫「交貨日」，宴客日叫「結婚日」。",
        orders=orders,
        notes=[
            "沒有單號：同一個分頁裡相鄰、姓名＋聯絡電話＋交貨日都相同的列，算同一張單；轉換工具替每張單補一個一看就知道是歷史資料的單號。",
            "品名混用品項別名（龍鳳、鳳梨酥、西式喜餅、西餅、太陽餅）。",
            "取貨日都已經過了，匯入後不會出現在廚房彙總。",
        ],
    )
    return wb, exp


# ---------------------------------------------------------------- entry point


def _digits(phone: str) -> str:
    d = "".join(ch for ch in phone if ch.isdigit())
    return "0" + d[3:] if d.startswith("886") else d


def generate(out_dir: Path, seed: int, today: date) -> Manifest:
    """Write every template and mock file under `out_dir` and return the 對照表."""
    out_dir = Path(out_dir)
    rng = random.Random(seed)
    people = People(rng)
    ctxs = {s.code: _Ctx(s, _STYLES[s.code], rng, people, today) for s in STORES}
    day1_stamp = datetime.combine(today, time(18, 0))
    day2 = today + timedelta(days=1)
    day2_stamp = datetime.combine(day2, time(18, 0))

    templates = []
    for store in STORES:
        rel = f"templates/{store.code}_{store.name}_喜餅訂單範本_{t.TEMPLATE_VERSION}.xlsx"
        t.write_template(out_dir / rel, store)
        templates.append(rel)

    def rel_day1(code: str) -> str:
        return f"day1/{code}_{store_by_code(code).name}.xlsx"

    tc01 = _tc01(ctxs["TC01"], rel_day1("TC01"))
    tc02 = _tc02(ctxs["TC02"], rel_day1("TC02"))
    tc03 = _tc03(ctxs["TC03"], rel_day1("TC03"))
    ch01 = _ch01(ctxs["CH01"], rel_day1("CH01"))
    ch02 = _ch02(ctxs["CH02"], rel_day1("CH02"))
    o110 = tc01.orders["0110"]
    nt01 = _nt01(ctxs["NT01"], rel_day1("NT01"), _digits(o110.phone), o110.orderer_name, o110.pickup_on)
    day1_results = [tc01, tc02, tc03, ch01, ch02, nt01]
    for result in day1_results:
        save_workbook(result.workbook, out_dir / result.expectation.file, day1_stamp)

    tc01_d2_rel = f"day2/TC01_{store_by_code('TC01').name}_修改.xlsx"
    tc01_d2_wb, tc01_d2 = _tc01_day2(ctxs["TC01"], tc01, tc01.expectation.file, tc01_d2_rel)
    save_workbook(tc01_d2_wb, out_dir / tc01_d2_rel, day2_stamp)

    nt01_resend_rel = f"day2/NT01_{store_by_code('NT01').name}_整份重傳.xlsx"
    (out_dir / nt01_resend_rel).write_bytes((out_dir / nt01.expectation.file).read_bytes())
    nt01_resend = _resend_expectation(nt01.expectation, store_by_code("NT01"), nt01_resend_rel, day2)

    tc02_returning = next(o for o in tc02.orders.values() if o.date_style != DATE_ROC)
    ch01_returning = list(ch01.orders.values())[3]
    legacy_stamp = datetime.combine(today - timedelta(days=30), time(12, 0))
    legacy_tc02_rel = f"legacy/TC02_{store_by_code('TC02').name}_舊訂單.xlsx"
    legacy_tc02_wb, legacy_tc02 = _legacy_tc02(ctxs["TC02"], legacy_tc02_rel, _digits(tc02_returning.phone))
    save_workbook(legacy_tc02_wb, out_dir / legacy_tc02_rel, legacy_stamp)
    legacy_ch01_rel = f"legacy/CH01_{store_by_code('CH01').name}_舊訂單.xlsx"
    legacy_ch01_wb, legacy_ch01 = _legacy_ch01(ctxs["CH01"], legacy_ch01_rel, _digits(ch01_returning.phone))
    save_workbook(legacy_ch01_wb, out_dir / legacy_ch01_rel, legacy_stamp)

    def seen(file: str, order: MockOrder) -> dict[str, str]:
        return {"file": file, "ref": f"單號 {order.order_no}", "orderer_name": order.orderer_name, "phone_as_typed": order.phone}

    def seen_legacy(lg: LegacyExpectation, o: LegacyOrder) -> dict[str, str]:
        return {
            "file": lg.file,
            "ref": f"{o.source_sheet} 第 {'、'.join(map(str, o.source_rows))} 列",
            "orderer_name": o.orderer_name,
            "phone_as_typed": o.phone,
        }

    nt01_o1 = next(iter(nt01.orders.values()))
    tc02_legacy_order = next(o for o in legacy_tc02.orders if _digits(o.phone) == _digits(tc02_returning.phone))
    ch01_legacy_order = next(o for o in legacy_ch01.orders if _digits(o.phone) == _digits(ch01_returning.phone))
    returning = [
        ReturningCustomer(
            _digits(tc01.orders["0101"].phone),
            [seen(tc01.expectation.file, tc01.orders["0101"]), seen(tc01.expectation.file, tc01.orders["0105"])],
            "同一位訂購人替兩對新人訂喜餅（同一支電話、兩張單）。",
        ),
        ReturningCustomer(
            _digits(o110.phone),
            [seen(tc01.expectation.file, o110), seen(nt01.expectation.file, nt01_o1)],
            "同一支電話在兩家門市下過單，寫法不同（+886），訂購人姓名也不同，要標出來。",
        ),
        ReturningCustomer(
            _digits(tc02_returning.phone),
            [seen_legacy(legacy_tc02, tc02_legacy_order), seen(tc02.expectation.file, tc02_returning)],
            "匯入歷史資料後才看得出來的回頭客。",
        ),
        ReturningCustomer(
            _digits(ch01_returning.phone),
            [seen_legacy(legacy_ch01, ch01_legacy_order), seen(ch01.expectation.file, ch01_returning)],
            "匯入歷史資料後才看得出來的回頭客。",
        ),
    ]

    manifest = Manifest(
        seed=seed,
        today=today.isoformat(),
        day2=day2.isoformat(),
        assumptions=dict(ASSUMPTIONS),
        templates=templates,
        files=[r.expectation for r in day1_results] + [tc01_d2, nt01_resend],
        legacy=[legacy_tc02, legacy_ch01],
        returning_customers=returning,
    )
    (out_dir / "expectations.json").write_text(manifest.to_json(), encoding="utf-8")
    (out_dir / "對照表.md").write_text(render_markdown(manifest), encoding="utf-8")
    return manifest
