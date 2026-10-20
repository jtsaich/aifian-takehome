"""convert-legacy: a 門市's pre-go-live Excel (歷史資料) turned into the 範本.

An FDE onboarding tool, not a product feature: the FDE runs it once per legacy file,

    uv run convert-legacy TC02 mock_data/legacy/TC02_台中北屯店_舊訂單.xlsx

and a 總部人員 imports the result on the 歷史資料匯入 page. Each 門市 kept its own
layout, so there is one converter per 門市 (`CONVERTERS`). Legacy orders have no
單號; each gets `H-<接單日 YYYYMMDD>-<NN>` (NN counts that 接單日's orders in file
order), which reads as 歷史資料 at a glance and never looks like a store's own
numbering. The output is byte-reproducible: converting the same file again gives
the same bytes, hence the same 單號, so re-importing never duplicates orders.

The converter only moves data into the 範本: 盒數 and 品項 are validated on
import like any upload. 品名 are written as the catalog's 品項 when they are a
name, 品項別名 or a known legacy short name, else as written (and listed, since
the import will refuse them). Anything it cannot place (a date without a year,
an unreadable 其他 cell, one order's rows disagreeing) stops the conversion.
"""

from __future__ import annotations

import argparse
import io
import re
import sys
import unicodedata
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from bakery import template as t
from bakery.catalog import PRODUCTS, store_by_code
from bakery.parsing import parse_date
from bakery.xlsx import workbook_bytes

ORDER_NO_PREFIX = "H-"
DEFAULT_OUT_DIR = "legacy_converted"

# Short names the legacy files use that are not 品項別名 the stores use today.
LEGACY_SHORT_NAMES: dict[str, str] = {"蛋黃酥": "蛋黃酥禮盒"}

_PRODUCT_NAMES: dict[str, str] = {
    **LEGACY_SHORT_NAMES,
    **{alias: p.name for p in PRODUCTS for alias in p.aliases},
    **{p.name: p.name for p in PRODUCTS},
}


class ConversionError(Exception):
    """The legacy file has something the converter can't place; `problems` say where."""

    def __init__(self, problems: Sequence[str]) -> None:
        super().__init__("\n".join(problems))
        self.problems = tuple(problems)


@dataclass
class _Line:
    product: str
    quantity: Any  # as written: the import validates it
    note: str | None = None


@dataclass
class _Order:
    received_on: date
    pickup_on: date
    orderer_name: str | None
    phone: str | None
    banquet_on: date | None = None
    couple_names: str | None = None
    lines: list[_Line] = field(default_factory=list)


@dataclass(frozen=True)
class Conversion:
    content: bytes  # the 範本 .xlsx
    order_count: int
    line_count: int
    unknown_products: tuple[str, ...]  # 品名 not in the catalog: the import will refuse those rows


class _Problems:
    def __init__(self) -> None:
        self.items: list[str] = []

    def add(self, ws: Worksheet, row: int, text: str) -> None:
        self.items.append(f"{ws.title} 第 {row} 列：{text}")


def convert(store_code: str, file_bytes: bytes) -> Conversion:
    """Convert `store_code`'s legacy file. Raises ConversionError, KeyError for a 門市 without a converter."""
    converter = CONVERTERS[store_code]
    problems = _Problems()
    orders = converter(load_workbook(io.BytesIO(file_bytes), data_only=True), problems)
    if problems.items:
        raise ConversionError(problems.items)
    return _template(store_code, orders)


# ---------------------------------------------------------------- TC02 台中北屯店

_TC02_OTHER_ITEM = re.compile(r"^(.+?)(\d+)盒?$")
_TC02_OTHER_SEPARATORS = re.compile(r"[、,，;；\s]+")


def _tc02(wb: Workbook, problems: _Problems) -> list[_Order]:
    """One row per order: 龍鳳／鳳梨酥／綠豆椪 each have a 盒數 column, other 品項 go in 其他
    as text like 太陽餅20盒; 民國年 dates; no 單號, 宴客日, 新人 or 接單人."""
    ws = wb.worksheets[0]
    cols = _headers(ws, ("接單日期", "客人", "電話", "其他", "取貨日", "備註"))
    products = {i: _PRODUCT_NAMES[h] for h, i in _header_row(ws).items() if h in _PRODUCT_NAMES}
    orders = []
    for r, values in _data_rows(ws):
        received = _date(values[cols["接單日期"]], "接單日", ws, r, problems)
        pickup = _date(values[cols["取貨日"]], "取貨日", ws, r, problems)
        lines = [_Line(name, values[i]) for i, name in products.items() if values[i] not in (None, "")]
        other = _text(values[cols["其他"]])
        for item in _TC02_OTHER_SEPARATORS.split(other) if other else ():
            if m := _TC02_OTHER_ITEM.match(item):
                lines.append(_Line(_product(m.group(1)), int(m.group(2))))
            else:
                problems.add(ws, r, f"其他欄寫「{other}」，看不出品項和盒數（例如「太陽餅20盒」）。")
        if not lines:
            problems.add(ws, r, "沒有任何品項的盒數。")
        if received is None or pickup is None or not lines:
            continue
        lines[0].note = _text(values[cols["備註"]])
        orders.append(
            _Order(received, pickup, _text(values[cols["客人"]]), _text(values[cols["電話"]]), lines=lines)
        )
    return orders


# ---------------------------------------------------------------- CH01 彰化中華店


def _ch01(wb: Workbook, problems: _Problems) -> list[_Order]:
    """One row per 品項, a sheet per month; adjacent rows of one sheet with the same
    姓名, 聯絡電話 and 交貨日 are one order. No 單號 or 接單人."""
    orders = []
    for ws in wb.worksheets:
        cols = _headers(ws, ("訂購日", "姓名", "聯絡電話", "品名", "數量", "交貨日", "結婚日", "新人"))
        key = None
        order: _Order | None = None
        first_row: tuple = ()
        for r, values in _data_rows(ws):
            row_key = tuple(values[cols[c]] for c in ("姓名", "聯絡電話", "交貨日"))
            if row_key != key:
                key, first_row, order = row_key, values, None
                received = _date(values[cols["訂購日"]], "接單日", ws, r, problems)
                pickup = _date(values[cols["交貨日"]], "取貨日", ws, r, problems)
                banquet_cell = values[cols["結婚日"]]
                banquet = _date(banquet_cell, "宴客日", ws, r, problems) if banquet_cell else None
                if received and pickup:
                    order = _Order(
                        received,
                        pickup,
                        _text(values[cols["姓名"]]),
                        _text(values[cols["聯絡電話"]]),
                        banquet_on=banquet,
                        couple_names=_text(values[cols["新人"]]),
                    )
                    orders.append(order)
            else:
                differ = [c for c in ("訂購日", "結婚日", "新人") if values[cols[c]] != first_row[cols[c]]]
                if differ:
                    problems.add(ws, r, f"跟上一列是同一張單，但{'、'.join(differ)}不一樣。")
            if order is not None:
                order.lines.append(_Line(_product(_text(values[cols["品名"]]) or ""), values[cols["數量"]]))
    return orders


CONVERTERS: dict[str, Callable[[Workbook, _Problems], list[_Order]]] = {"TC02": _tc02, "CH01": _ch01}


# ---------------------------------------------------------------- shared


def _header_row(ws: Worksheet) -> dict[str, int]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {_text(v): i for i, v in enumerate(first) if _text(v)}


def _headers(ws: Worksheet, needed: Sequence[str]) -> dict[str, int]:
    found = _header_row(ws)
    missing = [h for h in needed if h not in found]
    if missing:
        raise ConversionError([f"{ws.title}：第 1 列找不到欄位 {'、'.join(missing)}，這不是這家門市的舊格式。"])
    return found


def _data_rows(ws: Worksheet):
    for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if any(v not in (None, "") for v in values):
            yield r, values


# Two "today"s far apart: a date written with its year reads the same against both.
_REFERENCE_DAYS = (date(2000, 1, 1), date(2050, 7, 1))


def _date(value: Any, column: str, ws: Worksheet, row: int, problems: _Problems) -> date | None:
    """A date with its year; a year-less one (10/25) would be guessed, so it is a problem."""
    a, b = (parse_date(value, today) if value not in (None, "") else None for today in _REFERENCE_DAYS)
    if a is None or a != b:
        problems.add(ws, row, f"{column}「{_text(value) or ''}」看不懂，要有年月日。")
        return None
    return a


def _product(written: str) -> str:
    name = unicodedata.normalize("NFKC", written).strip()
    return _PRODUCT_NAMES.get(name, name)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _template(store_code: str, orders: Sequence[_Order]) -> Conversion:
    wb = t.build_template(store_by_code(store_code))
    ws = wb[t.SHEET_NAME]
    col = {name: i for i, name in enumerate(t.COLUMNS, start=1)}
    per_day: dict[date, int] = {}
    row = t.FIRST_DATA_ROW
    for order in orders:
        per_day[order.received_on] = n = per_day.get(order.received_on, 0) + 1
        order_no = f"{ORDER_NO_PREFIX}{order.received_on:%Y%m%d}-{n:02d}"
        for line in order.lines:
            values = {
                t.COL_ORDER_NO: order_no,
                t.COL_ORDERER_NAME: order.orderer_name,
                t.COL_PHONE: order.phone,
                t.COL_RECEIVED_ON: order.received_on,
                t.COL_PICKUP_ON: order.pickup_on,
                t.COL_BANQUET_ON: order.banquet_on,
                t.COL_COUPLE_NAMES: order.couple_names,
                t.COL_PRODUCT: line.product,
                t.COL_QUANTITY: line.quantity,
                t.COL_NOTE: line.note,
            }
            for name, value in values.items():
                cell = ws.cell(row=row, column=col[name], value=value)
                if name in t.DATE_COLUMNS:
                    cell.number_format = t.DATE_NUMBER_FORMAT
                elif name in t.TEXT_COLUMNS:
                    cell.number_format = "@"
            row += 1
    unknown = dict.fromkeys(
        line.product for o in orders for line in o.lines if line.product not in _PRODUCT_NAMES.values()
    )
    return Conversion(
        content=workbook_bytes(wb, t.TEMPLATE_TIMESTAMP),
        order_count=len(orders),
        line_count=sum(len(o.lines) for o in orders),
        unknown_products=tuple(unknown),
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="把門市上線前的舊 Excel 轉成喜餅訂單範本（FDE 導入工具）")
    parser.add_argument("store_code", choices=sorted(CONVERTERS), help="門市代碼：用這家門市的舊格式轉換")
    parser.add_argument("source", type=Path, help="舊格式 .xlsx")
    parser.add_argument(
        "-o", "--out", type=Path, default=None,
        help=f"輸出檔（預設 {DEFAULT_OUT_DIR}/<門市代碼>_<門市>_歷史資料.xlsx）",
    )
    args = parser.parse_args(argv)
    store = store_by_code(args.store_code)
    out = args.out or Path(DEFAULT_OUT_DIR) / f"{store.code}_{store.name}_歷史資料.xlsx"
    try:
        result = convert(store.code, args.source.read_bytes())
    except ConversionError as exc:
        print(f"沒有轉換：{args.source} 有 {len(exc.problems)} 個地方要先改：", file=sys.stderr)
        for problem in exc.problems:
            print(f"・{problem}", file=sys.stderr)
        return 1
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(result.content)
    print(f"{store.name}（{store.code}）：{result.order_count} 張喜餅訂單、{result.line_count} 筆訂單明細 → {out}")
    if result.unknown_products:
        print(f"注意，這些品名不在商品清單裡：{'、'.join(result.unknown_products)}。匯入時這幾列會被退，請先請總部加品項或別名。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
